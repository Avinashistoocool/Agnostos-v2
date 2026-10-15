"""Excel ingestion (Mode B) and result writing for both modes."""
from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Professor

# Canonical header -> Professor attribute. Matching is fuzzy/lowercased.
HEADER_MAP = {
    "professor name": "name",
    "name": "name",
    "professor": "name",
    "email address": "email",
    "email": "email",
    "contact email": "email",
    "institution": "institution",
    "university": "institution",
    "affiliation": "institution",
    "country": "country",
    "department": "department",
    "group": "department",
    "research area": "research_area",
    "research interests": "research_area",
    "research": "research_area",
    "fit with my profile": "fit_note",
    "fit": "fit_note",
    "funding status": "funding_status",
    "funding": "funding_status",
    "intl/indian eligibility": "eligibility",
    "eligibility": "eligibility",
    "application route": "application_route",
    "route": "application_route",
    "deadline": "deadline",
    "priority/rank": "rank",
    "rank": "rank",
    "priority": "rank",
    "probability": "prior_probability",
    "profile url": "profile_url",
    "url": "profile_url",
    "website": "profile_url",
    "notes": "notes",
    "openalex id": "openalex_id",
}

# Columns this tool writes/updates.
RESULT_COLUMNS = [
    "Local Match Score",
    "Probability of Acceptance",
    "Probability Justification",
    "Email Subject",
    "Custom Email Draft",
    "AI-Tell Score",
    "Review Flags",
    "Email Status",
    "Email Confidence",
    "Email Notes",
    "Suggested Email",
    "Profile Status",
    "LLM Status",
    "Send Status",
    "Sent At",
    "Reply Status",
    "Follow-up Notes",
]

OUTPUT_COLUMNS = [
    "Rank", "Professor Name", "Email Address", "Institution", "Country",
    "Department", "Research Area", "Profile URL", "Recent Publications",
    "Funding Status", "Eligibility", "Application Route", "Deadline",
    "OpenAlex ID",
] + RESULT_COLUMNS

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
RESULT_FILL = PatternFill("solid", fgColor="2E7D32")


def _norm(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def _find_header_row(ws, scan_rows: int = 8) -> int:
    """Return the 1-based row index that looks most like a header row."""
    best_row, best_hits = 1, -1
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(scan_rows, ws.max_row),
                                           values_only=True), start=1):
        hits = sum(1 for c in row if _norm(c) in HEADER_MAP)
        if hits > best_hits:
            best_row, best_hits = idx, hits
    return best_row


def _column_index(ws, header_row: int) -> dict[str, int]:
    """Map lowercased header text -> 1-based column index."""
    out = {}
    for col in range(1, ws.max_column + 1):
        key = _norm(ws.cell(row=header_row, column=col).value)
        if key:
            out[key] = col
    return out


def _is_blank_record(prof: Professor) -> bool:
    return not (prof.name or prof.email)


def load_targets(path: str | Path, sheet: str | None = None) -> list[Professor]:
    """Read professors from an existing workbook (Mode B).

    Reads every sheet that has a recognizable header unless `sheet` is given.
    De-duplicates by email (falling back to name+institution).
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Target workbook not found: {path}")
    wb = load_workbook(path, data_only=True)
    sheets = [wb[sheet]] if sheet else list(wb.worksheets)

    professors: list[Professor] = []
    seen: set[str] = set()
    for ws in sheets:
        header_row = _find_header_row(ws)
        headers = _column_index(ws, header_row)
        mapped = {col: HEADER_MAP[name] for name, col in headers.items() if name in HEADER_MAP}
        if "name" not in mapped.values():
            continue  # not a professor sheet (e.g. "Funding Programs")
        extra = {col: ws.cell(row=header_row, column=col).value
                 for col in headers.values() if col not in mapped}
        for r in range(header_row + 1, ws.max_row + 1):
            prof = Professor(source_sheet=ws.title, source_row=r)
            for col, attr in mapped.items():
                val = ws.cell(row=r, column=col).value
                if val is None:
                    continue
                if isinstance(val, datetime):
                    val = val.strftime("%Y-%m-%d")
                setattr(prof, attr, str(val).strip())
            notes = [f"{h}: {ws.cell(row=r, column=c).value}"
                     for c, h in extra.items()
                     if h and ws.cell(row=r, column=c).value
                     and str(h).strip() not in RESULT_COLUMNS]
            if notes:
                prof.notes = (prof.notes + " | " + " | ".join(notes)).strip(" |")
            if _is_blank_record(prof):
                continue
            key = prof.email.lower() or f"{prof.name.lower()}|{prof.institution.lower()}"
            if key in seen:
                continue
            seen.add(key)
            professors.append(prof)
    return professors


def backup(path: str | Path) -> Path:
    path = Path(path)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    dest = path.with_name(f"{path.stem}.backup-{stamp}{path.suffix}")
    shutil.copy2(path, dest)
    return dest


def _style_sheet(ws, header_row: int, result_cols: dict[str, int]) -> None:
    for col in result_cols.values():
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = RESULT_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = {"Probability Justification": 55, "Custom Email Draft": 80,
              "Email Subject": 42, "Local Match Score": 14,
              "Probability of Acceptance": 16, "Review Flags": 40,
              "Email Status": 12, "Email Confidence": 14, "Email Notes": 45,
              "Suggested Email": 32, "AI-Tell Score": 12}
    for name, col in result_cols.items():
        ws.column_dimensions[get_column_letter(col)].width = widths.get(name, 16)


def update_workbook(path: str | Path, professors: list[Professor],
                    make_backup: bool = True) -> Path:
    """Mode B: append/refresh result columns in-place on the source workbook."""
    path = Path(path)
    if make_backup:
        backup(path)
    wb = load_workbook(path)
    by_sheet: dict[str, list[Professor]] = {}
    for p in professors:
        if p.source_sheet and p.source_row:
            by_sheet.setdefault(p.source_sheet, []).append(p)

    for sheet_name, profs in by_sheet.items():
        ws = wb[sheet_name]
        header_row = _find_header_row(ws)
        headers = _column_index(ws, header_row)
        result_cols: dict[str, int] = {}
        next_col = ws.max_column + 1
        for name in RESULT_COLUMNS:
            existing = headers.get(name.lower())
            if existing:
                result_cols[name] = existing
            else:
                ws.cell(row=header_row, column=next_col, value=name)
                result_cols[name] = next_col
                next_col += 1
        _style_sheet(ws, header_row, result_cols)
        # Identity columns the user owns: only filled in when empty (recheck
        # can discover them), never overwritten.
        fill_cols = {attr: headers[h] for h, attr in
                     (("openalex id", "openalex_id"), ("institution", "institution"))
                     if h in headers}
        for p in profs:
            for attr, col in fill_cols.items():
                cell = ws.cell(row=p.source_row, column=col)
                if cell.value in (None, "") and getattr(p, attr, ""):
                    cell.value = getattr(p, attr)
            values = _result_values(p)
            for name, col in result_cols.items():
                val = values.get(name)
                if val in (None, ""):
                    continue  # never clobber manually-maintained cells
                cell = ws.cell(row=p.source_row, column=col, value=val)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return path


def _result_values(p: Professor) -> dict:
    return {
        "Local Match Score": round(p.local_score, 4) if p.local_score else None,
        "Probability of Acceptance": f"{p.probability}%" if p.probability else None,
        "Probability Justification": p.justification or None,
        "Email Subject": p.email_subject or None,
        "Custom Email Draft": p.email_body or None,
        "AI-Tell Score": p.ai_score if p.email_body else None,
        "Review Flags": p.ai_flags or None,
        "Email Status": p.email_status or None,
        "Email Confidence": f"{p.email_confidence}%" if p.email_status else None,
        "Email Notes": p.email_notes or None,
        "Suggested Email": p.email_suggestion or None,
        "Profile Status": p.profile_status or None,
        "LLM Status": p.llm_status or None,
    }


def _output_row(idx: int, p: Professor) -> list:
    res = _result_values(p)
    return [
        idx, p.name, p.email, p.institution, p.country, p.department,
        p.research_area, p.profile_url, "; ".join(p.recent_works[:4]),
        p.funding_status, p.eligibility, p.application_route, p.deadline,
        p.openalex_id,
        res["Local Match Score"], res["Probability of Acceptance"],
        res["Probability Justification"], res["Email Subject"],
        res["Custom Email Draft"], res["AI-Tell Score"], res["Review Flags"],
        res["Email Status"], res["Email Confidence"], res["Email Notes"],
        res["Suggested Email"], res["Profile Status"],
        p.llm_status, "", "", "", "",
    ]


def append_rows(path: str | Path, professors: list[Professor],
                sheet_name: str = "Targets") -> Path:
    """Add rows to a workbook this tool wrote, creating it if needed.

    Earlier rows are left untouched, so batch N never rewrites (or loses the
    verification results of) batches 1..N-1. Columns are matched by header
    name, so a workbook written by an older version still lines up.
    """
    path = Path(path)
    if not path.exists():
        return write_new_workbook(path, professors, sheet_name=sheet_name)
    wb = load_workbook(path)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]
    header_row = _find_header_row(ws)
    headers = _column_index(ws, header_row)
    next_col = ws.max_column + 1
    for name in OUTPUT_COLUMNS:
        if name.lower() not in headers:
            ws.cell(row=header_row, column=next_col, value=name)
            headers[name.lower()] = next_col
            next_col += 1
    last = header_row
    for r in range(ws.max_row, header_row, -1):
        if any(ws.cell(row=r, column=c).value not in (None, "")
               for c in range(1, min(ws.max_column, 4) + 1)):
            last = r
            break
    rank = last - header_row
    for p in professors:
        last += 1
        rank += 1
        for name, val in zip(OUTPUT_COLUMNS, _output_row(rank, p)):
            if val not in (None, ""):
                ws.cell(row=last, column=headers[name.lower()], value=val)
        p.source_sheet, p.source_row = ws.title, last
    wb.save(path)
    return path


def write_new_workbook(path: str | Path, professors: list[Professor],
                       sheet_name: str = "Targets") -> Path:
    """Mode A: write a fresh results workbook."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(OUTPUT_COLUMNS)
    for col in range(1, len(OUTPUT_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = RESULT_FILL if OUTPUT_COLUMNS[col - 1] in RESULT_COLUMNS else HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for idx, p in enumerate(sorted(professors, key=lambda x: (-x.probability, -x.local_score)), 1):
        ws.append(_output_row(idx, p))
        p.source_sheet, p.source_row = sheet_name, idx + 1

    widths = {"Custom Email Draft": 80, "Probability Justification": 55,
              "Research Area": 35, "Recent Publications": 45, "Email Subject": 42,
              "Review Flags": 40, "AI-Tell Score": 12, "Email Notes": 45,
              "Suggested Email": 32}
    for col, name in enumerate(OUTPUT_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(name, 18)
    ws.freeze_panes = "B2"
    wb.save(path)
    return path


def load_for_sending(path: str | Path, sheet: str | None = None) -> list[dict]:
    """Read finalized rows (with drafts) for the dispatch module."""
    path = Path(path)
    wb = load_workbook(path, data_only=True)
    sheets = [wb[sheet]] if sheet else list(wb.worksheets)
    rows = []
    for ws in sheets:
        header_row = _find_header_row(ws)
        headers = _column_index(ws, header_row)
        if "custom email draft" not in headers:
            continue
        for r in range(header_row + 1, ws.max_row + 1):
            def get(col_name):
                col = headers.get(col_name.lower())
                return ws.cell(row=r, column=col).value if col else None

            body = get("Custom Email Draft")
            email = get("Email Address") or get("Email")
            if not body or not email:
                continue
            prob_raw = str(get("Probability of Acceptance") or "").strip().rstrip("%")
            try:
                prob = int(float(prob_raw))
            except ValueError:
                prob = 0
            rows.append({
                "sheet": ws.title,
                "row": r,
                "name": get("Professor Name") or get("Name") or "",
                "email": str(email).strip(),
                "institution": get("Institution") or "",
                "country": str(get("Country") or "").strip(),
                "subject": get("Email Subject") or "",
                "body": str(body),
                "probability": prob,
                "send_status": str(get("Send Status") or "").strip(),
                "sent_at": str(get("Sent At") or "").strip(),
                "message_id": str(get("Message ID") or "").strip(),
                "reply_status": str(get("Reply Status") or "").strip(),
                "email_status": str(get("Email Status") or "").strip().lower(),
                "suggested_email": str(get("Suggested Email") or "").strip(),
            })
    _apply_pending(path, rows)
    return rows


def _apply_pending(path: Path, rows: list[dict]) -> None:
    """Overlay statuses parked while the workbook was locked."""
    sidecar = sidecar_path(path)
    if not sidecar.exists():
        return
    try:
        pending = json.loads(sidecar.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return
    by_cell = {(p["sheet"], p["row"]): p for p in pending}
    for row in rows:
        hit = by_cell.get((row["sheet"], row["row"]))
        if hit:
            row["send_status"] = hit["status"]


def sidecar_path(path: str | Path) -> Path:
    """Where statuses are parked when the workbook cannot be written."""
    path = Path(path)
    return path.with_suffix(path.suffix + ".pending-status.json")


def is_writable(path: str | Path) -> bool:
    """True if the workbook is not held open by Excel or another process."""
    path = Path(path)
    if not path.exists():
        return False
    try:
        with open(path, "r+b"):
            return True
    except OSError:
        return False


def mark_sent(path: str | Path, updates: list[dict]) -> Path | None:
    """Write Send Status / Sent At back to the workbook.

    If the workbook is open in Excel the write fails, and after a long live
    send that would silently discard every status. Park the updates in a
    sidecar JSON instead and merge them on the next successful write, so a
    locked file can never cause a professor to be mailed twice.

    Returns the sidecar path if the write was deferred, else None.
    """
    path = Path(path)
    sidecar = sidecar_path(path)
    pending: list[dict] = []
    if sidecar.exists():
        try:
            pending = json.loads(sidecar.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            pending = []
    # Later updates for the same cell win.
    merged: dict[tuple, dict] = {}
    for upd in [*pending, *updates]:
        merged[(upd["sheet"], upd["row"])] = upd
    updates = list(merged.values())

    try:
        wb = load_workbook(path)
    except OSError:
        _park(sidecar, updates)
        return sidecar
    for upd in updates:
        ws = wb[upd["sheet"]]
        header_row = _find_header_row(ws)
        headers = _column_index(ws, header_row)
        names = ["Send Status", "Sent At"] + (["Message ID"] if upd.get("message_id") else [])
        for name in names:
            if name.lower() not in headers:
                col = ws.max_column + 1
                ws.cell(row=header_row, column=col, value=name)
                headers[name.lower()] = col
        ws.cell(row=upd["row"], column=headers["send status"], value=upd["status"])
        ws.cell(row=upd["row"], column=headers["sent at"], value=upd.get("sent_at", ""))
        if upd.get("message_id"):
            ws.cell(row=upd["row"], column=headers["message id"], value=upd["message_id"])
    try:
        wb.save(path)
    except OSError:
        _park(sidecar, updates)
        return sidecar
    sidecar.unlink(missing_ok=True)
    return None


def _park(sidecar: Path, updates: list[dict]) -> None:
    slim = [{"sheet": u["sheet"], "row": u["row"], "status": u["status"],
             "sent_at": u.get("sent_at", ""),
             "message_id": u.get("message_id", "")} for u in updates]
    sidecar.write_text(json.dumps(slim, indent=2), encoding="utf-8")

def mark_replies(path: str | Path, replies: dict) -> None:
    """Write Reply Status / Follow-up Notes for {(sheet, row): Reply}."""
    path = Path(path)
    wb = load_workbook(path)
    for (sheet, row), reply in replies.items():
        ws = wb[sheet]
        header_row = _find_header_row(ws)
        headers = _column_index(ws, header_row)
        for name in ("Reply Status", "Follow-up Notes"):
            if name.lower() not in headers:
                col = ws.max_column + 1
                ws.cell(row=header_row, column=col, value=name)
                headers[name.lower()] = col
        ws.cell(row=row, column=headers["reply status"], value=reply.status())
        note = "[reply] " + reply.subject + (f" | {reply.snippet}" if reply.snippet else "")
        cell = ws.cell(row=row, column=headers["follow-up notes"])
        # Never clobber notes you typed yourself.
        if not cell.value or str(cell.value).startswith("[reply]"):
            cell.value = note[:500]
    wb.save(path)
