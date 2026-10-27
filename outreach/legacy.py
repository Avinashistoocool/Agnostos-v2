"""`legacy`: find eminent researchers in your field, and write them a different
kind of email - an aspiring scientist's humble, specific letter, not an
internship form letter.

Who counts as "legacy":
  * very high lifetime impact (h-index / citations) in YOUR topics, from the
    OpenAlex author index (worldwide, not just Europe), and still active, plus
  * anyone you name yourself (`--include "Neil deGrasse Tyson; Fabiola Gianotti"`).

Every candidate is paired with a real summer route they could take you through
(CERN Summer Student Programme, DESY, Fermilab SULI/IPS, Perimeter, KITP,
Simons, their group's visitor scheme...). The LLM must only name a programme
from that list, or ask them which route they'd suggest.

Files (all in output/, git-ignored):
  legacy.xlsx         candidates, address, status, draft, sent
Sending uses the RESEND_SMTP_* account (your personal Gmail).
"""
from __future__ import annotations

import json
import re
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .discover import Discoverer, resolve_countries

COLUMNS = ["Name", "Email", "Institution", "Country", "h-index", "Citations",
           "Known for", "Recent work", "Programme route", "Why them", "Source",
           "Email Status", "Email Notes", "Subject", "Draft", "Approved",
           "Send Status", "Sent At", "Message ID", "OpenAlex ID"]

# Summer routes a non-EU undergraduate can actually apply to. The LLM may only
# name one of these (or ask the researcher which they'd suggest).
PROGRAMMES = {
    "cern": "CERN Summer Student Programme (non-member state stream; applications ~Dec-Jan)",
    "desy": "DESY Summer Student Programme (Hamburg/Zeuthen; applications ~Jan)",
    "fermilab": "Fermilab International Physics/SIST-style summer programmes",
    "slac": "SLAC summer research programmes",
    "psi": "PSI Summer Student Programme",
    "gsi": "GSI/FAIR Summer Student Programme",
    "perimeter": "Perimeter Institute visiting / undergraduate research programmes",
    "kitp": "KITP (UC Santa Barbara) programmes",
    "ictp": "ICTP (Trieste) schools and visiting programmes",
    "mpi": "Max Planck Institute internships / IMPRS summer schools",
    "daad": "DAAD WISE (German labs, for Indian students; applications ~Nov-Dec)",
    "mitacs": "Mitacs Globalink Research Internship (Canada, for Indian students)",
    "charpak": "Charpak Lab Scholarship (France, for Indian students)",
    "epfl": "EPFL Excellence Research Internship Program",
    "eth": "ETH Zurich Student Summer Research Fellowship",
    "caltech": "Caltech SURF (international applicants via a host professor)",
    "nasa": "NASA / JPL visiting student programmes",
    "group": "a short visiting-student stay in their own group",
}

INST_ROUTES = [
    (r"cern|european organization for nuclear", "cern"), (r"desy|deutsches elektronen", "desy"),
    (r"fermi", "fermilab"), (r"slac|stanford linear", "slac"), (r"paul scherrer", "psi"),
    (r"gsi|helmholtzzentrum f.r schwerionen", "gsi"), (r"perimeter", "perimeter"),
    (r"kavli institute for theoretical|santa barbara", "kitp"), (r"ictp|abdus salam", "ictp"),
    (r"max planck", "mpi"), (r"epfl|lausanne", "epfl"), (r"eth z", "eth"),
    (r"caltech|california institute of technology", "caltech"),
    (r"nasa|jet propulsion", "nasa"),
]
COUNTRY_ROUTES = {"DE": "daad", "CA": "mitacs", "FR": "charpak", "CH": "cern"}


def programme_for(institution: str, country: str) -> str:
    inst = (institution or "").lower()
    keys = [k for pat, k in INST_ROUTES if re.search(pat, inst)]
    if not keys and country in COUNTRY_ROUTES:
        keys.append(COUNTRY_ROUTES[country])
    keys.append("group")
    return "; ".join(PROGRAMMES[k] for k in dict.fromkeys(keys))


@dataclass
class Legend:
    name: str
    institution: str = ""
    country: str = ""
    h_index: int = 0
    citations: int = 0
    topics: list[str] = field(default_factory=list)
    recent: list[str] = field(default_factory=list)
    openalex_id: str = ""
    source: str = "openalex"
    email: str = ""
    email_status: str = ""
    email_notes: str = ""


# ----------------------------------------------------------------- finding
def _author_rows(disc: Discoverer, params: dict) -> list[dict]:
    data = disc._api("/authors", params) or {}
    return data.get("results") or []


def search_legends(disc: Discoverer, topics: list[str], countries: list[str],
                   min_h: int, per_topic: int, active_since: int) -> list[Legend]:
    """Highest-impact authors whose work matches each topic and who still publish."""
    found: dict[str, Legend] = {}
    for topic in topics:
        # Works search ranks by relevance; collect the authors of the most-cited
        # relevant papers, then rank those authors by lifetime impact.
        flt = ["cited_by_count:>200", f"from_publication_date:{active_since - 10}-01-01"]
        if countries:
            flt.append("authorships.countries:" + "|".join(countries))
        works = (disc._api("/works", {"search": topic, "filter": ",".join(flt),
                                     "sort": "cited_by_count:desc", "per-page": 100,
                                     "select": "authorships"}) or {}).get("results") or []
        ids = []
        for w in works:
            for a in (w.get("authorships") or [])[:12]:
                aid = ((a.get("author") or {}).get("id") or "").rsplit("/", 1)[-1]
                if aid and aid not in ids:
                    ids.append(aid)
        rows = []
        for start in range(0, min(len(ids), 200), 50):
            rows += _author_rows(disc, {
                "filter": "openalex:" + "|".join(ids[start:start + 50]),
                "per-page": 50,
                "select": "id,display_name,summary_stats,cited_by_count,"
                          "last_known_institutions,topics,counts_by_year"})
        picked = 0
        for r in sorted(rows, key=lambda r: -((r.get("summary_stats") or {}).get("h_index") or 0)):
            h = (r.get("summary_stats") or {}).get("h_index") or 0
            if h < min_h:
                break
            recent = sum(c.get("works_count", 0) for c in (r.get("counts_by_year") or [])
                         if c.get("year", 0) >= active_since)
            if recent < 2:               # retired / no longer active
                continue
            inst = (r.get("last_known_institutions") or [{}])[0] or {}
            aid = r["id"]
            if aid in found:
                continue
            found[aid] = Legend(
                name=r.get("display_name", ""), institution=inst.get("display_name", ""),
                country=(inst.get("country_code") or "").upper(), h_index=h,
                citations=r.get("cited_by_count") or 0,
                topics=[t.get("display_name", "") for t in (r.get("topics") or [])[:5]],
                openalex_id=aid, source=f"top impact in '{topic}'")
            picked += 1
            if picked >= per_topic:
                break
    return list(found.values())


def lookup_named(disc: Discoverer, names: list[str]) -> list[Legend]:
    out = []
    for name in names:
        rows = _author_rows(disc, {"search": name, "per-page": 5,
                                   "select": "id,display_name,summary_stats,cited_by_count,"
                                             "last_known_institutions,topics"})
        if not rows:
            out.append(Legend(name=name, source="you named them (not in OpenAlex)"))
            continue
        r = max(rows, key=lambda r: r.get("cited_by_count") or 0)
        inst = (r.get("last_known_institutions") or [{}])[0] or {}
        out.append(Legend(
            name=r.get("display_name", name), institution=inst.get("display_name", ""),
            country=(inst.get("country_code") or "").upper(),
            h_index=(r.get("summary_stats") or {}).get("h_index") or 0,
            citations=r.get("cited_by_count") or 0,
            topics=[t.get("display_name", "") for t in (r.get("topics") or [])[:5]],
            openalex_id=r["id"], source="you named them"))
    return out


def add_recent_work(disc: Discoverer, legends: list[Legend]) -> None:
    for lg in legends:
        if not lg.openalex_id:
            continue
        aid = lg.openalex_id.rsplit("/", 1)[-1]
        data = disc._api("/works", {"filter": f"author.id:{aid}", "sort": "publication_date:desc",
                                    "per-page": 25,
                                    "select": "title,publication_year,authorships"}) or {}
        # Only papers they plausibly shaped: big-collaboration papers list
        # thousands of names and say nothing about this person's own work.
        lg.recent = [f"{w.get('title')} ({w.get('publication_year')})"
                     for w in data.get("results") or []
                     if w.get("title") and len(w.get("authorships") or []) <= 12][:3]


def is_physics(lg: "Legend") -> bool:
    """Stricter than focus.in_focus: at least two distinct physics terms, no off-field topics."""
    from .focus import OFF_FOCUS, focus_pattern
    text = " ".join(lg.topics + lg.recent)
    hits = {m.group(0).lower() for m in focus_pattern("").finditer(text)}
    return len(hits) >= 2 and not re.search(OFF_FOCUS + r"|material|graphene|battery|climate|"
                                            r"carbon|nutrition|child|health", text, re.I)


_US_LIKE = {"US", "", None}


def address_fits(email: str, institution: str, country: str) -> bool:
    """Reject addresses at an unrelated institution (kwatanab@fgcu.edu for NIMS Japan)."""
    from .verify import country_mismatch, domain_of
    dom = domain_of(email)
    if country_mismatch(country, email):
        return False
    if dom.endswith(".edu") and country not in _US_LIKE:
        return False
    # Domain should share a word with the institution, or be a known lab/uni pattern.
    inst_words = {w for w in re.findall(r"[a-z]{3,}", (institution or "").lower())}
    dom_words = set(re.findall(r"[a-z]{3,}", dom.lower()))
    initials = "".join(w[0] for w in re.findall(r"[A-Za-z]+", institution or "")
                       if w[0].isupper()).lower()
    if not institution:
        from .verify import is_academic
        return is_academic(dom)            # unknown employer: only accept academic domains
    if inst_words & dom_words:
        return True
    if len(initials) >= 2 and initials in dom.lower():
        return True
    # Abbreviated domains: ox.ac.uk -> Oxford, wisc.edu -> Wisconsin, cam -> Cambridge.
    labels = [l for l in dom.lower().split(".") if len(l) >= 2]
    return any(w.startswith(l) for l in labels for w in inst_words
               if l not in {"ac", "edu", "co", "uni", "www", "mail", "physics", "phys", "astro"})


def find_emails(legends: list[Legend], workers: int = 4) -> None:
    """Same finder as `verify --repair`, with the strict wrong-person checks."""
    from .models import Professor
    from .verify import verify_many
    profs = [Professor(name=lg.name, institution=lg.institution, country=lg.country)
             for lg in legends]
    from .config import settings
    from .verify import load_cache, save_cache
    cache_path = settings.output_dir / "legacy_email_cache.json"
    cache = load_cache(cache_path)
    verify_many(profs, workers=workers, repair=True, apply=True, cache=cache)
    save_cache(cache_path, cache)
    for lg, p in zip(legends, profs):
        lg.email = p.email if p.email_status == "ok" else ""
        if lg.email and not address_fits(lg.email, lg.institution, lg.country):
            p.email_notes = f"rejected {lg.email}: not at {lg.institution}"
            lg.email, p.email_status = "", "invalid"
        lg.email_status = p.email_status or "invalid"
        lg.email_notes = (p.email_notes or "")[:200] + (
            "" if lg.email else " | public figure: use their official contact page")


# ----------------------------------------------------------------- workbook
def load(path: Path) -> list[dict]:
    if not path.exists():
        return []
    ws = load_workbook(path).active
    hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, [c.value for c in row])) for row in ws.iter_rows(min_row=2)
            if any(c.value for c in row)]


def save(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Legacy researchers"
    ws.append(COLUMNS)
    for r in rows:
        ws.append([r.get(c, "") for c in COLUMNS])
    for col, w in zip("ABCDEFGHIJKLMNOPQRST", (24, 32, 30, 7, 8, 10, 40, 50, 50, 40, 22, 10,
                                              40, 40, 80, 9, 10, 18, 30, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def merge(rows: list[dict], legends: list[Legend]) -> list[dict]:
    have = {(r.get("OpenAlex ID") or r.get("Name") or "").lower() for r in rows}
    for lg in legends:
        key = (lg.openalex_id or lg.name).lower()
        if key in have:
            continue
        have.add(key)
        rows.append({
            "Name": lg.name, "Email": lg.email, "Institution": lg.institution,
            "Country": lg.country, "h-index": lg.h_index, "Citations": lg.citations,
            "Known for": "; ".join(lg.topics), "Recent work": "; ".join(lg.recent),
            "Programme route": programme_for(lg.institution, lg.country),
            "Source": lg.source, "Email Status": lg.email_status,
            "Email Notes": lg.email_notes, "Approved": "", "Send Status": "",
            "OpenAlex ID": lg.openalex_id})
    rows.sort(key=lambda r: -(int(r.get("h-index") or 0)))
    return rows


# ----------------------------------------------------------------- drafting
LEGACY_SYSTEM = """You help an undergraduate physics student write ONE letter to an eminent scientist they have admired for years. This is NOT an internship form letter. It is an aspiring scientist's humble, sincere, specific note, the kind a senior researcher remembers: short, warm, honest, never flattering in generic terms.

APPLICANT PROFILE
{profile}

Availability: {availability}
Funding: {funding}

THE SCIENTIST
{who}

Real summer routes this scientist could plausibly support (name ONLY from this list, or ask which route they'd suggest):
{routes}

HOW TO WRITE IT
- 160-220 words. Plain, personal, a little vulnerable; no grandiose adjectives ("legendary", "iconic", "visionary"), no "I am writing to", no em dashes, no markdown, no bullet lists.
- Start with "Dear Professor <Surname>," (or "Dear Dr. <Surname>," if not a professor) on its own line.
- Then ONE concrete thing of theirs that genuinely shaped how the student thinks. Use ONLY "Recent work" titles (those are their own small-team papers) or the broad "Known for" fields. Never attribute a specific paper, model or survey that is not in those lines; if in doubt, speak about their field of work in general terms.
- Then who the student is, in two sentences, tied to that idea through ONE real project from the profile (numbers only if they are in the profile).
- The ask, small and humble: whether they might be open to the student spending summer {summer} working under them or in their group through one of the routes above, or, if that isn't possible, whether they could point the student to a person or programme. Say honestly they understand the scientist's time is limited and a short reply, or none, is completely fine.
- One closing line that sounds like a person, then "Warm regards," and the student's first name. The signature is added automatically.
- Never claim to have met them, read something not listed, or have skills not in the profile. Never invent personal history (when or how the student first came across their work, a talk they attended, a book they read as a child). If something essential is unknown, write [FILL: ...].

Return only JSON: {{"subject": "<short, human subject, no clickbait>", "body": "<email body>"}}"""


def who_block(r: dict) -> str:
    return "\n".join(x for x in (
        f"Name: {r.get('Name')}",
        f"Institution: {r.get('Institution')} ({r.get('Country')})",
        f"Known for: {r.get('Known for')}",
        f"Recent work: {r.get('Recent work')}",
        f"h-index {r.get('h-index')}, {r.get('Citations')} citations") if x)


def draft(evaluator, settings, profile, r: dict, summer: str) -> tuple[str, str]:
    from .humanize import deslop
    system = LEGACY_SYSTEM.format(profile=profile.summary_block(),
                                  availability=settings.availability,
                                  funding=settings.funding_stance, who=who_block(r),
                                  routes=r.get("Programme route") or PROGRAMMES["group"],
                                  summer=summer)
    data = evaluator.chat_json(system, "Write the letter.")
    return (str(data.get("subject", "")).strip()[:120],
            deslop(str(data.get("body", "")).strip()))
