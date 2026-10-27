"""CLI entrypoint for the academic outreach pipeline.

  python -m outreach.cli profile
  python -m outreach.cli discover --count 1000 --batches 4 --query "topic a; topic b"
  python -m outreach.cli recheck --xlsx targets.xlsx --apply
  python -m outreach.cli run  --mode list   [--yes] [--no-enrich] [--max-professors N]
  python -m outreach.cli run  --mode scrape --sources sources.yaml
  python -m outreach.cli verify --repair
  python -m outreach.cli send --dry-run
  python -m outreach.cli send --live
  python -m outreach.cli stats
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

from .config import MODEL_PRICES, PROVIDERS, ConfigError, settings
from .batches import plan_batches, run_batches, split_queries
from .discover import COUNTRY_GROUPS
from .enrich import OpenAlexClient
from .excel_io import (_column_index, _find_header_row, append_rows, is_writable,
                       load_for_sending, load_targets, mark_replies, mark_sent,
                       sidecar_path, update_workbook, write_new_workbook)
from .llm import LLMEvaluator, confirm_gate, list_models
from .mailer import Mailer
from .matcher import filter_candidates
from .resume_parser import parse_resume, save_profile
from .recheck import recheck_all
from .replies import ReplyChecker
from .schedule import SendWindow, fmt_duration, parse_at
from .stats import append_send_log, collect, read_send_log
from .verify import load_cache, save_cache, verify_many

BAR = "=" * 62

# Mapped from the exception name the mailer reports, so a failed run ends with
# the fix rather than a bare traceback class.
SMTP_HINTS = {
    "SMTPSenderRefused":
        "   The server refused your From address. SMTP_FROM usually has to be\n"
        "   exactly the mailbox SMTP_USER authenticates as - check .env.",
    "SMTPAuthenticationError":
        "   Login rejected. For Gmail use a 16-char App Password, not your\n"
        "   account password (https://myaccount.google.com/apppasswords).",
    "SMTPServerDisconnected":
        "   The server dropped the connection, usually rate-limiting. Raise\n"
        "   SEND_DELAY_MIN/MAX, lower DAILY_SEND_CAP, and re-run to continue.",
    "SMTPRecipientsRefused":
        "   The recipient address was rejected - check that row in the sheet.",
    "SMTPDataError":
        "   The server rejected the message body, often as suspected spam.\n"
        "   Shorten the draft or send fewer at a time.",
    "SMTPConnectError":
        "   Could not reach SMTP_HOST:SMTP_PORT. Check the host, port, and\n"
        "   SMTP_SECURITY (starttls on 587, ssl on 465).",
    "TimeoutError":
        "   The mail server timed out. Raise SMTP_TIMEOUT in .env.",
}


def _h(title: str) -> None:
    print(f"\n{BAR}\n {title}\n{BAR}")


def _safe(text: str) -> str:
    """Windows consoles are often cp1252; never let printing kill a run."""
    enc = sys.stdout.encoding or "utf-8"
    return str(text).encode(enc, errors="replace").decode(enc)


def _load_sources(path: Path) -> list[dict]:
    if not path.exists():
        raise FileNotFoundError(f"Sources file not found: {path}")
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    else:
        try:
            import yaml
        except ImportError as exc:
            raise RuntimeError("pip install pyyaml, or use a .json sources file") from exc
        data = yaml.safe_load(path.read_text(encoding="utf-8-sig"))
    sources = data.get("sources", data) if isinstance(data, dict) else data
    if not isinstance(sources, list):
        raise ValueError("Sources file must contain a list of {url, institution, country}")
    return sources


# ---------------------------------------------------------------- commands
def cmd_profile(args) -> int:
    _h("RESUME PARSE")
    profile = parse_resume(settings.require_resume(args.resume))
    print(_safe(f" Name        : {profile.name}"))
    print(_safe(f" Email       : {profile.email}"))
    print(_safe(f" Programme   : {profile.degree} @ {profile.institution} "
                f"({profile.graduation_year}, GPA {profile.gpa})"))
    print(_safe(f" Interests   : {', '.join(profile.research_interests)}"))
    print(_safe(f" Skills      : {', '.join(profile.skills[:20])}"))
    print(f" Projects    : {len(profile.projects)}")
    for p in profile.projects:
        print(_safe(f"   - {p['title']}"))
    settings.ensure_dirs()
    out = settings.output_dir / "profile.json"
    save_profile(profile, out)
    print(f"\n Saved structured profile -> {out}")
    print("\n--- LLM system-prompt profile block ---")
    print(_safe(profile.summary_block()))
    return 0


SKIPPABLE = ("enrich", "filter", "gate")


def _parse_skip(values: list[str]) -> set[str]:
    steps = {s.strip().lower() for v in values for s in v.split(",") if s.strip()}
    bad = steps - set(SKIPPABLE)
    if bad:
        raise SystemExit(f"--skip: unknown step(s) {', '.join(sorted(bad))}; "
                         f"choose from {', '.join(SKIPPABLE)}")
    return steps


def _select_without_filter(professors, max_count: int):
    selected, rejected = [], []
    for p in professors:
        if not p.email:
            p.llm_status = "skipped: no email"
        elif len(selected) >= max_count:
            p.llm_status = "skipped: over max-professors cap"
        else:
            selected.append(p)
            continue
        rejected.append(p)
    return selected, rejected


def cmd_run(args) -> int:
    skip = _parse_skip(args.skip)
    if args.no_enrich:
        skip.add("enrich")
    if "gate" in skip:
        args.yes = True
    settings.ensure_dirs()
    if args.provider:
        settings.provider = args.provider
        settings.draft_model = args.model or PROVIDERS[args.provider]["default_model"]
    if args.max_professors:
        settings.max_professors = args.max_professors
    if args.budget:
        settings.budget_usd = args.budget
    if args.model:
        settings.draft_model = args.model

    _h("STEP 1/6  RESUME")
    profile = parse_resume(settings.require_resume(args.resume))
    print(_safe(f" {profile.name} | {profile.degree} @ {profile.institution}"))
    print(_safe(f" Interests: {', '.join(profile.research_interests[:6])}"))

    _h("STEP 2/6  TARGET ACQUISITION")
    if args.mode == "scrape":
        from .scraper import run_scrape
        sources = _load_sources(Path(args.sources))
        print(f" Scraping {len(sources)} directory page(s) with stealth browser...")

        def scrape_progress(i, total, src, result, found):
            flag = "BLOCKED" if result.blocked else ("ERR" if result.error else "ok")
            print(_safe(f"  [{i}/{total}] {src['url'][:62]:<62} {flag:>8} "
                        f"+{len(found)} profs"))
            if result.error:
                print(_safe(f"          {result.error[:90]}"))

        professors = run_scrape(settings, sources, scrape_progress,
                                resolve_profiles=not args.no_profile_hop)
        with_email = sum(1 for p in professors if p.email)
        print(f" Discovered {len(professors)} people, {with_email} with an email address.")
    else:
        xlsx = settings.require_target_xlsx(args.xlsx)
        professors = load_targets(xlsx, sheet=args.sheet)
        print(f" Loaded {len(professors)} professors from {xlsx.name}")
        # Rewriting a sent row's draft would lose the record of what was
        # actually mailed (and re-sends/follow-ups must use that text).
        professors = _drop_sent(xlsx, professors)
        if settings.focus and not args.no_focus:
            from .focus import in_focus
            before = len(professors)
            professors = [p for p in professors if in_focus(p, settings.focus_keywords)]
            print(f" Focus (physics/HEP): keeping {len(professors)} of {before} rows; "
                  "set FOCUS=false or pass --no-focus to draft everyone.")
    if not professors:
        print(" No professors found. Aborting.")
        return 1

    _h("STEP 3/6  OPENALEX ENRICHMENT")
    if "enrich" in skip or not settings.use_openalex:
        print(" Skipped.")
    else:
        client = OpenAlexClient(
            mailto=settings.openalex_mailto or profile.email,
            api_key=settings.openalex_api_key,
            cache_path=settings.output_dir / "openalex_cache.json",
            works_per_author=settings.openalex_works,
        )

        def enrich_progress(i, total, p):
            if i % 10 == 0 or i == total:
                print(f"  [{i}/{total}] enriched")

        client.enrich_all(professors, enrich_progress)
        hit = sum(1 for p in professors if p.recent_works)
        print(f" Publication data found for {hit}/{len(professors)} professors.")
        if client.rate_limited:
            print(f" ! {client.last_error}")
            print("   Continuing with directory/sheet text only.")

    _h("STEP 4/6  LOCAL NLP FILTER (zero tokens)")
    min_score = args.min_score if args.min_score is not None else settings.min_local_score
    if "filter" in skip:
        print(" Skipped: every row with an email goes to the LLM, in sheet order.")
        selected, rejected = _select_without_filter(professors, settings.max_professors)
    else:
        selected, rejected = filter_candidates(
            professors, profile, min_score=min_score, max_count=settings.max_professors,
        )
    print(f" {len(selected)} selected / {len(rejected)} filtered out "
          f"(threshold {min_score}, cap {settings.max_professors})")
    reasons = Counter(
        "low local score" if p.llm_status.startswith("skipped: local score")
        else p.llm_status.removeprefix("skipped: ")
        for p in rejected
    )
    for reason, n in reasons.most_common():
        print(_safe(f"   {n:>5} x {reason}"))
    for p in selected[:12]:
        print(_safe(f"   {p.local_score:>6.3f}  {p.name[:30]:<30} {p.research_area[:36]}"))
    if len(selected) > 12:
        print(f"   ... and {len(selected) - 12} more")
    if not selected:
        if reasons.get("no email", 0) == len(rejected):
            print(" Nothing passed: no row has an email address yet. Fill them first:")
            print(f"   python -m outreach.cli verify --xlsx \"{args.xlsx or settings.target_xlsx}\" "
                  f"--repair --apply")
        elif reasons.get("low local score"):
            print(" Nothing passed the filter. Lower --min-score and retry.")
        else:
            print(" Nothing passed the filter (see reasons above).")
        return 1

    _h("STEP 5/6  LLM EVALUATION + DRAFTING")
    cost_note = " (free tier)" if settings.is_free_tier else ""
    print(f" Provider: {settings.provider}{cost_note} | model: {settings.draft_model}")
    settings.require_api_key()
    evaluator = LLMEvaluator(settings, profile)
    estimate = evaluator.estimate(selected)
    if not confirm_gate(estimate, settings.budget_usd, assume_yes=args.yes):
        print(" Aborted by user. No API calls made, no files written.")
        return 2

    def llm_progress(i, total, p):
        status = f"{p.probability:>3}%" if p.llm_status.startswith("ok") else "----"
        print(_safe(f"  [{i}/{total}] {status}  {p.name[:34]:<34} {p.llm_status[:28]}"))

    evaluator.evaluate_all(selected, progress=llm_progress,
                           budget_usd=settings.budget_usd)
    print("\n " + evaluator.usage_report())

    _h("STEP 6/6  WRITE RESULTS")
    all_profs = selected + rejected
    if args.mode == "scrape" or args.output:
        out_path = Path(args.output) if args.output else (
            settings.output_dir / f"outreach_targets_{datetime.now():%Y%m%d-%H%M}.xlsx")
        write_new_workbook(out_path, all_profs)
        print(f" Wrote new workbook -> {out_path}")
    else:
        target = settings.require_target_xlsx(args.xlsx)
        update_workbook(target, all_profs, make_backup=not args.no_backup)
        print(f" Updated in place    -> {target}")
        if not args.no_backup:
            print(" Timestamped backup saved alongside the original.")

    ranked = sorted([p for p in selected if p.probability and p.llm_status.startswith("ok")],
                    key=lambda p: -p.probability)
    if ranked:
        print("\n Top matches:")
        for p in ranked[:10]:
            print(_safe(f"   {p.probability:>3}%  {p.name[:32]:<32} {p.institution[:26]}"))
    else:
        print("\n No drafts were produced. Check the LLM Status column for errors.")
        return 1
    print("\n Next: review the drafts, then run "
          "`python -m outreach.cli send --dry-run`.")
    return 0


def _focus_text(path) -> dict:
    """(sheet,row) -> fields used by focus.in_focus."""
    out = {}
    for p in load_targets(path):
        out[(p.source_sheet, p.source_row)] = {
            "research_area": p.research_area, "institution": p.institution,
            "department": getattr(p, "department", ""),
            "recent_works": " ".join(p.recent_works or [])}
    return out


def cmd_send(args) -> int:
    settings.ensure_dirs()
    if args.delay_min is not None:
        settings.send_delay_min = args.delay_min
    if args.delay_max is not None:
        settings.send_delay_max = args.delay_max
    if args.min_probability is not None:
        settings.send_min_probability = args.min_probability
    if args.limit:
        settings.daily_send_cap = args.limit

    path = settings.require_target_xlsx(args.xlsx)
    profile = parse_resume(settings.require_resume(args.resume))
    rows = load_for_sending(path, sheet=args.sheet)
    if settings.focus and not args.no_focus:
        from .focus import in_focus
        focus_text = _focus_text(path)
        before = len(rows)
        rows = [r for r in rows if r["send_status"].lower() == "sent"
                or in_focus(focus_text.get((r["sheet"], r["row"]), {}), settings.focus_keywords)]
        if before != len(rows):
            print(f" Focus (physics/HEP): holding back {before - len(rows)} off-topic draft(s).")

    window = None
    window_spec = args.local_window if args.local_window is not None else settings.send_local_window
    if window_spec and window_spec.lower() not in {"off", "none", "any"}:
        try:
            window = SendWindow(window_spec, weekends=args.weekends)
        except ValueError as exc:
            print(f" ERROR: {exc}")
            return 1
    start_at = None
    if args.at:
        try:
            start_at = parse_at(args.at)
        except ValueError:
            print(f" ERROR: --at wants HH:MM or 'YYYY-MM-DD HH:MM', got {args.at!r}")
            return 1

    # --- Export mode: write .eml files you can open in any mail client ---
    if args.export:
        out_dir = Path(args.export_dir) if args.export_dir else settings.output_dir / "drafts"
        _h(f"EXPORT DRAFTS -> {out_dir}")
        print(f" {len(rows)} rows with drafts in {path.name}")
        if not rows:
            print(" Nothing to export. Run the `run` command first.")
            return 1
        mailer = Mailer(settings, profile, dry_run=True)
        eligible = [r for r in rows
                    if r["probability"] >= settings.send_min_probability]
        eligible.sort(key=lambda r: -r["probability"])
        if args.top:
            eligible = eligible[:args.top]
        for r in eligible:
            written = mailer.save_eml(r, out_dir)
            print(_safe(f"  {r['probability']:>3}%  {written.name}"))
        print(f"\n Wrote {len(eligible)} .eml files to {out_dir}")
        print(" Double-click any of them to open in Outlook/Thunderbird/Mail.")
        return 0

    review_to = args.review_to or (settings.smtp_from if args.review else "")
    dry_run = not (args.live or review_to)
    mode = "DRY RUN" if dry_run else ("REVIEW -> " + review_to if review_to else "LIVE")
    _h(f"EMAIL DISPATCH ({mode})")
    print(f" {len(rows)} rows with drafts in {path.name}")
    if not rows:
        print(" Nothing to send. Run the `run` command first.")
        return 1

    if not dry_run and not is_writable(path):
        print(_safe(f" ERROR: {path.name} is open in Excel (or otherwise locked)."))
        print(" Close it before sending — statuses must be written back, or the")
        print(" same professors could be mailed twice on the next run.")
        return 1

    if not dry_run:
        target_desc = (f"ALL to {review_to} (review copies)" if review_to
                       else "the professors themselves")
        print(f" Sending from {settings.smtp_from} via "
              f"{settings.smtp_host}:{settings.smtp_port} ({settings.smtp_security})")
        print(f" Recipients: {target_desc}")
        print(f" Delay {settings.send_delay_min:.0f}-{settings.send_delay_max:.0f}s | "
              f"cap {settings.daily_send_cap} | min probability "
              f"{settings.send_min_probability}%")
        if args.top:
            print(f" Mailing only the {args.top} highest-probability rows.")
        if settings.attachments:
            print(f" Attachments: {', '.join(Path(a).name for a in settings.attachments)}")
        if start_at:
            print(f" Scheduled start: {start_at:%a %d %b %Y %H:%M} (your local time)")
        if window:
            print(f" Local window: {window.spec} recipient time"
                  f"{'' if window.weekends else ', weekdays only'}"
                  f"{' - waiting for it' if args.wait else ' - others held for a later run'}")
        prompt = (" Send review copies to yourself now? [Y/N]: " if review_to
                  else " Send REAL emails to professors now? [Y/N]: ")
        if args.yes:
            print(" --yes given: not asking for confirmation.")
        elif input(prompt).strip().lower() not in {"y", "yes"}:
            print(" Aborted.")
            return 2
    elif window:
        print(f" Local window: {window.spec} recipient time (dry run ignores --at/--wait)")

    if start_at and not dry_run:
        wait = (start_at - datetime.now()).total_seconds()
        if wait > 0:
            print(" Leave this window open (and the PC awake) until it finishes.")
            Mailer._countdown(wait, label=f"sending starts at {start_at:%H:%M}")

    def progress(i, total, row, result):
        if result["status"] == "aborted":
            print(_safe(f"\n  ! Stopped: {result['detail']}"))
            return
        if result["status"] == "held":
            held[0] += 1
            if held[0] <= 5:
                print(_safe(f"  [{i}/{total}] held         {row['email'][:40]:<40} "
                            f"({row.get('country') or '?'}: {result['detail']})"))
            return
        print(_safe(f"  [{i}/{total}] {result['status']:<12} {row['email'][:40]:<40} "
                    f"({row['probability']}%)"))
        if result["status"].startswith("failed") and result.get("detail"):
            print(_safe(f"       -> {result['detail'][:120]}"))
        if dry_run and i <= args.preview and result.get("message") is not None:
            msg = result["message"]
            print("  " + "-" * 58)
            print(_safe(f"  Subject: {msg['Subject']}"))
            for line in msg.get_body(("plain",)).get_content().splitlines():
                print(_safe("  | " + line))
            print("  " + "-" * 58)

    held = [0]

    def checkpoint(update):
        # Persist each real delivery immediately: a crash or closed window
        # mid-run must not forget who was already mailed.
        if update["status"] in {"sent", "review-sent"}:
            mark_sent(path, [update])

    with Mailer(settings, profile, dry_run=dry_run, redirect_to=review_to) as mailer:
        updates = mailer.send_batch(rows, progress=progress, resend=args.resend,
                                    top=args.top,
                                    allow_unverified=args.allow_unverified,
                                    window=window, wait_for_window=args.wait,
                                    on_update=None if dry_run else checkpoint)
    if held[0]:
        print(f"\n {held[0]} row(s) held: outside their local {window.spec} window. "
              "Re-run later, or add --wait to sit until each window opens.")

    deferred = mark_sent(path, updates)
    try:
        append_send_log(settings.output_dir / "send_log.csv",
                        "dry-run" if dry_run else ("review" if review_to else "live"),
                        rows, updates)
    except OSError as exc:
        print(_safe(f" (could not append send_log.csv: {exc})"))
    sent = sum(1 for u in updates if u["status"] == "sent")
    review = sum(1 for u in updates if u["status"] == "review-sent")
    failed = [u for u in updates if u["status"].startswith("failed")]
    print(f"\n Done. sent={sent} review-sent={review} "
          f"dry-run={sum(1 for u in updates if u['status'] == 'dry-run')} "
          f"failed={len(failed)} | statuses written back to {path.name}")
    if deferred:
        print(_safe(f" NOTE: {path.name} was locked, so statuses were parked in "
                    f"{deferred.name}."))
        print(" They are applied automatically on the next run — nobody will be mailed twice.")
    if failed:
        first = failed[0]
        print(_safe(f" First failure: {first['status']}"))
        if first.get("detail"):
            print(_safe(f"   {first['detail'][:200]}"))
        print(SMTP_HINTS.get(first["status"].removeprefix("failed: "),
                             "   Re-run the same command to retry the failed rows; "
                             "successful sends are not repeated."))
    if dry_run:
        print("\n To actually see these in your mail app, either:")
        print("   python -m outreach.cli send --review    # mail copies to yourself")
        print("   python -m outreach.cli send --export    # write .eml files to open")
    return 0


def cmd_stats(args) -> int:
    path = settings.require_target_xlsx(args.xlsx)
    data = collect(path, sheet=args.sheet)
    c = data["counts"]
    _h(f"OUTREACH STATS - {path.name}")

    def section(title, counter):
        if not counter:
            return
        print(f"\n {title}")
        for k, v in counter.most_common():
            print(_safe(f"   {v:>5}  {k}"))

    total = c["total"] or 1
    print(" Funnel")
    for label, n in [("professors in sheet", c["total"]),
                     ("with an email address", c["with_email"]),
                     ("drafted by the LLM", c["drafted"]),
                     ("sent to the professor", c["sent"]),
                     ("review copies sent to you", c["review_sent"])]:
        print(f"   {n:>5}  {label:<28} ({100 * n / total:.0f}%)")
    replied = data["reply"].get("replied", 0)
    if c["sent"]:
        print(f"   {replied:>5}  {'replied':<28} ({100 * replied / c['sent']:.0f}% of sent)")

    section("Email verification", data["email"])
    section("Why rows were / weren't drafted (LLM Status)", data["llm"])
    section("Send Status", data["send"])
    section("Replies", data["reply"])
    section("Sent per day", Counter(dict(sorted(data["sent_days"].items()))))
    if data["pending_sidecar"]:
        print(f"\n {data['pending_sidecar']} status update(s) parked in the sidecar "
              "(workbook was locked); counted above.")

    log = read_send_log(settings.output_dir / "send_log.csv")
    if log:
        print("\n Send log (all attempts, incl. overwritten statuses)")
        for (mode, status), n in sorted(log.items()):
            print(_safe(f"   {n:>5}  {mode:<8} {status}"))

    # Resends from the second account (`resend`) never touch Send Status.
    resent_today = resent_all = 0
    log_path = settings.output_dir / "send_log.csv"
    if log_path.exists():
        import csv
        today = f"{datetime.now():%Y-%m-%d}"
        with log_path.open(newline="", encoding="utf-8") as fh:
            for rec in csv.DictReader(fh):
                if rec.get("mode") == "resend" and rec.get("status") == "sent":
                    resent_all += 1
                    resent_today += rec.get("time", "").startswith(today)
    if resent_all:
        print(f"\n Resends (original email again from {settings.resend_smtp_from or 'second account'})")
        print(f"   {resent_all:>5}  total   ({resent_today} today)")

    # Follow-up replies live in reply_queue.json, not the workbook.
    queue_path = settings.output_dir / "reply_queue.json"
    if queue_path.exists():
        try:
            q = json.loads(queue_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            q = []
        kinds = Counter()
        for item in q:
            kind = item.get("kind") or "reply"
            if item.get("action") == "thank":
                kind = "thank-you"
            state = "sent" if item.get("sent_at") else (
                "skipped (nothing to say)" if item.get("action") == "none" else
                "approved, waiting" if item.get("approved") else "drafted, needs approval")
            kinds[f"{kind:<10} {state}"] += 1
        section("Replies & follow-ups we sent (reply_queue.json)", kinds)

    replies_sent = 0
    if queue_path.exists():
        replies_sent = sum(1 for i in q if i.get("sent_at"))
    print(f"\n Total emails delivered: {c['sent'] + resent_all + replies_sent} "
          f"({c['sent']} first emails + {resent_all} resends + {replies_sent} replies)")
    print(" (send_log.csv only started on 19 Oct, so its 'live' count is lower than "
          "the workbook's; the workbook is the full record.)")

    llm = data["llm"]
    hints = []
    if llm.get("skipped: over max-professors cap"):
        hints.append(f"{llm['skipped: over max-professors cap']} rows were never drafted "
                     f"because of MAX_PROFESSORS={settings.max_professors}. Re-run "
                     "`run --max-professors N` to draft more.")
    if llm.get("skipped: no email"):
        hints.append(f"{llm['skipped: no email']} rows have no address; try `verify --repair`.")
    if llm.get("skipped: local score below threshold"):
        hints.append("Some rows fell below MIN_LOCAL_SCORE; lower it with `run --min-score`.")
    waiting = data["send"].get("not attempted yet", 0) + data["send"].get("dry-run", 0)
    if waiting:
        hints.append(f"{waiting} drafted rows not yet mailed (dry-run only or untouched); "
                     f"DAILY_SEND_CAP={settings.daily_send_cap} limits each live run.")
    if hints:
        print("\n Why so few were sent")
        for h in hints:
            print(_safe(f"   - {h}"))
    return 0


def _reply_countries(args) -> dict[str, str]:
    """email -> Country column from the workbook, so replies use the same timezone."""
    try:
        path = settings.require_target_xlsx(args.xlsx)
        return {r["email"].lower(): r.get("country", "")
                for r in load_for_sending(path, sheet=args.sheet) if r.get("email")}
    except Exception:
        return {}


def _tld_country(email: str) -> str:
    tld = email.rsplit(".", 1)[-1].lower()
    if tld == "uk":
        return "GB"
    return tld.upper() if len(tld) == 2 else ""


def cmd_resend(args) -> int:
    """Re-deliver the ORIGINAL email (no redraft) from a second account to
    professors who never answered, and/or whose first mail bounced."""
    import csv
    from dataclasses import replace
    from datetime import timedelta
    from openpyxl import load_workbook

    settings.ensure_dirs()
    path = settings.require_target_xlsx(args.xlsx)
    profile = parse_resume(settings.require_resume(args.resume))
    log_path = settings.output_dir / "send_log.csv"
    s = settings
    alt = replace(
        s, smtp_host=s.resend_smtp_host, smtp_port=s.resend_smtp_port,
        smtp_security=s.resend_smtp_security, smtp_user=s.resend_smtp_user,
        smtp_password=s.resend_smtp_password, smtp_from=s.resend_smtp_from,
        reply_to="", daily_send_cap=args.limit or s.resend_daily_cap,
        send_delay_min=args.delay_min if args.delay_min is not None else s.send_delay_min,
        send_delay_max=args.delay_max if args.delay_max is not None else s.send_delay_max)
    _h(f"RESEND ORIGINAL EMAILS FROM {alt.smtp_from or '?'} "
       f"({'LIVE' if args.live else 'DRY RUN'})")
    if args.live and not (alt.smtp_user and alt.smtp_password):
        print(" ERROR: set RESEND_SMTP_USER and RESEND_SMTP_PASSWORD (an app password) in .env")
        return 1

    # Everyone already resent - from the log (authoritative) and the workbook.
    done = set()
    today = 0
    if log_path.exists():
        with log_path.open(newline="", encoding="utf-8") as fh:
            for rec in csv.DictReader(fh):
                if rec.get("mode") == "resend" and rec.get("status") == "sent":
                    done.add(rec.get("email", "").strip().lower())
                    if rec.get("time", "").startswith(f"{datetime.now():%Y-%m-%d}"):
                        today += 1

    wb = load_workbook(path, read_only=True, data_only=True)
    resent_cells = set()
    for ws in wb.worksheets:
        hdr = _find_header_row(ws)
        col = _column_index(ws, hdr).get("resent at")
        if not col:
            continue
        for r, row in enumerate(ws.iter_rows(min_row=hdr + 1, values_only=True), hdr + 1):
            if col - 1 < len(row) and row[col - 1]:
                resent_cells.add((ws.title, r))
    wb.close()

    cutoff = datetime.now() - timedelta(days=args.min_days if args.min_days is not None
                                        else s.resend_min_days)
    want = {"bounced": args.who in ("all", "bounced"),
            "none": args.who in ("all", "noreply")}
    rows, skipped = [], Counter()
    for r in load_for_sending(path, sheet=args.sheet):
        if r["send_status"].lower() != "sent":
            continue
        kind = (r["reply_status"].split() or ["none"])[0].lower()
        if kind in ("replied", "auto-reply"):
            skipped["replied / auto-reply"] += 1
            continue
        if r["email_status"] == "invalid":
            # Hard bounce / wrong person: sending the same text again can't work.
            skipped["address known bad (see FAILURES.json)"] += 1
            continue
        if not want.get(kind, False):
            skipped[f"not selected ({kind})"] += 1
            continue
        if r["email"].lower() in done or (r["sheet"], r["row"]) in resent_cells:
            skipped["already resent"] += 1
            continue
        try:
            first = datetime.strptime(r["sent_at"][:19], "%Y-%m-%d %H:%M:%S")
        except ValueError:
            first = None
        if first and first > cutoff:
            skipped["first email too recent"] += 1
            continue
        r["send_status"] = ""
        rows.append(r)
    # Top universities first (Imperial, Oxford, ETH...), bounced first within a tier.
    ranks = {}
    if args.prioritize_top:
        from .discover import Discoverer, resolve_countries
        disc = Discoverer(mailto=s.openalex_mailto, api_key=s.openalex_api_key,
                          cache_path=s.output_dir / "discover_cache.json")
        unis = disc.top_universities(resolve_countries(args.region or s.discover_countries),
                                     args.prioritize_top, s.discover_since)
        disc.save_cache()
        ranks = {u["name"].strip().lower(): i for i, u in enumerate(unis, 1)}
        print(_safe(f" Prioritising top {len(ranks)} universities: "
                    + ", ".join(u["name"] for u in unis[:8]) + " ..."))
    rows.sort(key=lambda r: (ranks.get(str(r["institution"]).strip().lower(), 10_000),
                             not r["reply_status"].lower().startswith("bounced"),
                             r["sent_at"]))
    room = max(0, alt.daily_send_cap - today) if args.live else alt.daily_send_cap
    print(f" Eligible: {len(rows)} | already resent today: {today} | this run: "
          f"up to {min(room, len(rows))}")
    for why, n in skipped.most_common():
        print(_safe(f"   skipped {n:>4} x {why}"))
    if not rows or not room:
        return 0
    alt.daily_send_cap = room

    window = None
    spec = args.local_window if args.local_window is not None else s.send_local_window
    if spec and spec.lower() not in {"off", "none", "any"}:
        window = SendWindow(spec, weekends=args.weekends)
        for r in rows:
            r["country"] = r.get("country") or _tld_country(r["email"])
        print(f" Local window: {window.spec} recipient time"
              f"{' - waiting for it' if args.wait else ' - others held for a later run'}")
    if args.live and not args.yes and input(
            f" Resend up to {min(room, len(rows))} original emails from {alt.smtp_from}? "
            "[Y/N]: ").strip().lower() not in {"y", "yes"}:
        print(" Aborted.")
        return 2

    def record(update):
        if update["status"] != "sent":
            return
        src = next((r for r in rows if (r["sheet"], r["row"]) == (update["sheet"], update["row"])), {})
        # Log first: it is what blocks a second resend even if the workbook is locked.
        append_send_log(log_path, "resend", rows, [update])
        try:
            wb = load_workbook(path)
            ws = wb[update["sheet"]]
            hdr = _find_header_row(ws)
            cols = _column_index(ws, hdr)
            for name in ("Resent At", "Resent From", "Resend Message ID"):
                if name.lower() not in cols:
                    cols[name.lower()] = ws.max_column + 1
                    ws.cell(row=hdr, column=cols[name.lower()], value=name)
            ws.cell(row=update["row"], column=cols["resent at"], value=update["sent_at"])
            ws.cell(row=update["row"], column=cols["resent from"], value=alt.smtp_from)
            ws.cell(row=update["row"], column=cols["resend message id"],
                    value=update.get("message_id", ""))
            wb.save(path)
        except OSError:
            print(_safe(f"   ({path.name} locked; resend of {src.get('email')} is in "
                        "send_log.csv and will not repeat)"))

    def progress(i, total, row, result):
        if result["status"] == "held":
            return
        print(_safe(f"  [{i}/{total}] {result['status']:<10} {row['email'][:40]:<40} "
                    f"{str(row['institution'])[:30]:<30} "
                    f"({(row['reply_status'] or 'no reply')[:22]})"))
        if result["status"].startswith("failed"):
            print(_safe(f"       -> {result.get('detail', '')[:140]}"))

    with Mailer(alt, profile, dry_run=not args.live) as mailer:
        updates = mailer.send_batch(rows, progress=progress, resend=True,
                                    window=window, wait_for_window=args.wait,
                                    on_update=record if args.live else None)
    held = len(rows) - len(updates)
    c = Counter(u["status"] for u in updates)
    print(f"\n Done. {dict(c)}" + (f" | {held} not reached (cap/window)" if held else ""))
    if not args.live:
        print(" Dry run. Add --live to send.")
    return 0


def _qkey(q: dict) -> tuple:
    """Identity of a queued email: a follow-up shares in_reply_to with the reply before it."""
    return (q["email"], q.get("in_reply_to", ""), q.get("kind", "reply"), q.get("n", 0))


def _meet_dur():
    from datetime import timedelta
    return timedelta(minutes=settings.meet_minutes)


def _gcal_or_none():
    from . import gcal
    if not gcal.available():
        return None
    try:
        return gcal.service()
    except Exception as exc:
        print(_safe(f" (Google Calendar unavailable: {exc}; slots won't be checked "
                    "against your calendar)"))
        return None


def _held_slots(items) -> list:
    """Slots already proposed in unsent/sent queue items, so two professors aren't offered the same time."""
    out = []
    for q in items:
        for s in q.get("meet_slots") or []:
            try:
                start = datetime.fromisoformat(s)
            except ValueError:
                continue
            out.append((start, start + _meet_dur()))
    return out


def _calendar_holds(q: dict) -> list[str]:
    if not q.get("meet_slots"):
        return []
    from . import gcal
    svc = _gcal_or_none()
    if svc is None:
        print(_safe("             (calendar holds skipped: run `calendar-auth` first)"))
        return []
    try:
        slots = [datetime.fromisoformat(s) for s in q["meet_slots"]]
        ids = gcal.create_holds(svc, slots, q.get("name", ""), q["email"],
                                q.get("institution", ""))
        print(_safe(f"             + {len(ids)} calendar hold(s) added"))
        return ids
    except Exception as exc:
        print(_safe(f"             (calendar holds failed: {exc})"))
        return []


def _draft_funding_followups(threads, old, evaluator, profile, gcal_svc, taken,
                             args, queue_path) -> list:
    """Nudge professors who said they can't fund and then went quiet.

    Up to FUNDING_FOLLOWUPS extra emails per thread, each only after
    FOLLOWUP_GAP_DAYS of silence since our last message, each with a meeting
    request. Stops for good as soon as the professor writes again (that turns
    into a normal reply) or the limit is reached."""
    from datetime import timedelta
    from .replier import draft_followup, mentions_no_funding, plan_meeting, save_queue
    from . import gcal
    from .schedule import country_tz
    limit, gap = settings.funding_followups, timedelta(days=settings.followup_gap_days)
    made = 0
    for t in threads:
        if t.awaiting_us() or not any(
                m.sender == "professor" and mentions_no_funding(m.text) for m in t.messages):
            continue
        last_prof = t.last_incoming
        mine = [m for m in t.messages if m.sender == "student"
                and (m.when or datetime.min) > (last_prof.when or datetime.min)]
        if not mine:
            continue
        # Our follow-ups to this professor message, sent or still queued.
        ours = [q for q in old.values() if q["email"] == t.row["email"]
                and q.get("kind") == "followup" and q.get("in_reply_to") == last_prof.message_id]
        if any(not q.get("sent_at") for q in ours):
            continue                      # one is already waiting to go out
        n = len(ours) + 1
        if n > limit:
            continue
        last_mine = max((m.when for m in mine if m.when), default=None)
        if last_mine and datetime.now() - last_mine < gap:
            continue
        tz = country_tz(t.row.get("country")) or country_tz(t.row["email"].rsplit(".", 1)[-1])
        slots = gcal.propose_slots(tz, n=3, svc=gcal_svc, taken=taken)
        lines = gcal.describe(slots, tz)
        try:
            body = draft_followup(evaluator, t, settings, profile, lines)
        except Exception as exc:
            print(_safe(f"  ! follow-up {t.row['email']}: {exc}"))
            continue
        refs = " ".join(m.message_id for m in t.messages if m.message_id)
        item = {"email": t.row["email"], "name": t.row.get("name", ""),
                "institution": t.row.get("institution", ""),
                "subject": last_prof.subject or t.row.get("subject", ""),
                "in_reply_to": last_prof.message_id, "references": refs,
                "their_message": last_prof.text[:1500],
                "summary": f"Funding follow-up {n}/{limit} after {settings.followup_gap_days:g} "
                           "days of silence, with meeting request",
                "action": "followup", "body": body, "kind": "followup", "n": n,
                "meet_slots": [s.isoformat() for s in slots],
                "approved": bool(args.approve_all),
                "drafted_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "sent_at": ""}
        old[_qkey(item)] = item
        taken += [(s, s + _meet_dur()) for s in slots]
        save_queue(queue_path, list(old.values()))
        made += 1
        print(_safe(f"\n  [follow-up {n}/{limit}] {t.row.get('name')} <{t.row['email']}>"))
        for line in body.splitlines():
            print(_safe("    | " + line))
    if made:
        print(f"\n {made} funding follow-up(s) drafted.")
    return list(old.values())


def _merge_sent_queue(threads, queue) -> None:
    """Replies we sent over SMTP often have no copy in the IMAP Sent folder;
    add them to each thread so 'who wrote last' and message counts are right."""
    from .replier import Message
    by_email = {}
    for q in queue:
        if q.get("sent_at") and q.get("body"):
            by_email.setdefault(q["email"].lower(), []).append(q)
    for t in threads:
        for q in by_email.get(t.row["email"].lower(), []):
            when = datetime.strptime(q["sent_at"][:19], "%Y-%m-%d %H:%M:%S")
            if any(m.sender == "student" and m.when and abs((m.when - when).total_seconds()) < 600
                   for m in t.messages):
                continue
            t.messages.append(Message(when=when, sender="student",
                                      subject="Re: " + q.get("subject", ""), text=q["body"]))
        t.messages.sort(key=lambda m: m.when or datetime.min)


def cmd_calendar_auth(args) -> int:
    from . import gcal
    _h("GOOGLE CALENDAR AUTHORISATION")
    secret, token = gcal._paths()
    if not secret.exists() and not token.exists():
        print(f" Missing OAuth client file: {secret}")
        print(" See README > Google Calendar for the 2-minute setup.")
        return 1
    svc = gcal.service(interactive=True)
    cal = svc.calendars().get(calendarId=settings.google_calendar_id).execute()
    print(_safe(f" Authorised. Calendar: {cal.get('summary')} ({cal.get('timeZone')})"))
    return 0


def cmd_archive_sent(args) -> int:
    """Backfill the IMAP Sent folder with every email we sent over SMTP."""
    from email.utils import format_datetime, make_msgid
    from .mailbox import SentArchiver
    from .replier import load_queue
    settings.ensure_dirs()
    path = settings.require_target_xlsx(args.xlsx)
    profile = parse_resume(settings.require_resume(args.resume))
    mailer = Mailer(settings, profile, dry_run=True)
    _h("COPY SENT EMAILS INTO YOUR SENT FOLDER")
    items = []
    for r in load_for_sending(path):
        if r["send_status"].lower() != "sent":
            continue
        try:
            when = datetime.strptime(r["sent_at"][:19], "%Y-%m-%d %H:%M:%S")
        except ValueError:
            when = None
        msg = mailer.build_message(r["email"], r.get("name", ""),
                                   r.get("subject") or "Summer research internship enquiry",
                                   r["body"])
        items.append((msg, when, r.get("message_id", ""), "first"))
    for q in load_queue(settings.output_dir / "reply_queue.json"):
        if not q.get("sent_at") or not q.get("body"):
            continue
        when = datetime.strptime(q["sent_at"][:19], "%Y-%m-%d %H:%M:%S")
        try:
            msg = mailer.build_reply(q["email"], q.get("name", ""), q["subject"], q["body"],
                                     q.get("in_reply_to", ""), q.get("references", ""))
        except ValueError:
            continue
        items.append((msg, when, q.get("message_id", ""), q.get("kind", "reply")))
    # Remember what was copied: replies (and a few early sends) have no stored
    # Message-ID, so the IMAP duplicate check alone can't catch a re-run.
    done_path = settings.output_dir / "archived_sent.json"
    done = set(json.loads(done_path.read_text(encoding="utf-8"))) if done_path.exists() else set()
    def akey(msg, when):
        from email.utils import parseaddr
        return f"{parseaddr(str(msg['To']))[1].lower()}|{when:%Y-%m-%d %H:%M:%S}" if when else ""
    items = [it for it in items if akey(it[0], it[1]) not in done]
    print(f" {len(items)} sent emails to copy (first emails + replies); "
          f"{len(done)} copied on earlier runs.")
    if not args.live:
        print(" Dry run. Add --live to copy them into your Sent folder.")
        return 0
    added = skipped = failed = 0
    with SentArchiver(settings) as arc:
        print(_safe(f" Mailbox {arc.user} @ {arc.host}, folder '{arc.folder}'"))
        for i, (msg, when, mid, kind) in enumerate(items, 1):
            # Keep the original Message-ID so replies thread under it.
            if mid:
                del msg["Message-ID"]
                msg["Message-ID"] = mid
            if when:
                del msg["Date"]
                msg["Date"] = format_datetime(when.astimezone())
            # Attachments were sent but aren't needed to read the thread.
            if msg.is_multipart():
                plain = msg.get_body(("plain",)).get_content()
                for h in ("Content-Type", "Content-Transfer-Encoding", "MIME-Version"):
                    del msg[h]
                msg.set_content(plain)
            try:
                if arc.append(msg, when):
                    added += 1
                else:
                    skipped += 1
                if akey(msg, when):
                    done.add(akey(msg, when))
            except Exception as exc:
                failed += 1
                print(_safe(f"  ! {msg['To']}: {exc}"))
            if i % 25 == 0:
                print(f"  {i}/{len(items)}  added={added} already there={skipped}")
    done_path.write_text(json.dumps(sorted(done), indent=0), encoding="utf-8")
    print(f" Done. added={added} already there={skipped} failed={failed}")
    return 0


def cmd_conversations(args) -> int:
    """Write output/conversations.html: every thread, both sides, in one page."""
    from .convoview import build_html
    settings.ensure_dirs()
    path = settings.require_target_xlsx(args.xlsx)
    out = settings.output_dir / "conversations.html"
    _h("BUILD CONVERSATIONS PAGE")
    n = build_html(settings, path, out, progress=lambda f, i, t, r: (
        sys.stdout.write(f"\r  scanning {f}: {i}/{t}   "), sys.stdout.flush()))
    print(f"\n Wrote {n} conversations -> {out}")
    if not args.no_open:
        import webbrowser
        webbrowser.open(out.resolve().as_uri())
    return 0


def cmd_replace_failed(args) -> int:
    """Swap unreachable / wrong-person professors for a colleague in the same department."""
    from .replace import replace_failed
    settings.ensure_dirs()
    path = settings.require_target_xlsx(args.xlsx)
    fp = settings.output_dir / "FAILURES.json"
    _h("REPLACE UNREACHABLE PROFESSORS WITH A COLLEAGUE")
    if not fp.exists():
        print(" output/FAILURES.json not found; run `failures` first.")
        return 1
    if not is_writable(path):
        print(_safe(f" ERROR: {path.name} is open in Excel. Close it and re-run."))
        return 1
    res = replace_failed(settings, path, fp, limit=args.per_failure)
    print(f"\n Added {res['added']} colleague(s). Next: verify --repair --apply --only-missing, "
          "run --mode list, send.")
    return 0


def cmd_legacy(args) -> int:
    """Eminent researchers in your field + an aspiring-scientist letter, sent from your Gmail."""
    from dataclasses import replace
    from email.utils import make_msgid
    from . import legacy as L
    from .discover import Discoverer, resolve_countries
    settings.ensure_dirs()
    path = settings.output_dir / "legacy.xlsx"
    rows = L.load(path)
    profile = parse_resume(settings.require_resume(args.resume))

    if args.find or not rows:
        _h("LEGACY RESEARCHERS: FIND")
        topics = split_queries(args.topics or "particle physics; high energy physics phenomenology; "
                               "machine learning for particle physics; cosmology; "
                               "quantum field theory; astrophysics")
        disc = Discoverer(mailto=settings.openalex_mailto, api_key=settings.openalex_api_key,
                          cache_path=settings.output_dir / "discover_cache.json")
        legends = L.search_legends(disc, topics, resolve_countries(args.countries or ""),
                                   args.min_h, args.per_topic, datetime.now().year - 2)
        named = split_queries(args.include or "")
        if named:
            legends += L.lookup_named(disc, named)
        have = {(r.get("OpenAlex ID") or r.get("Name") or "").lower() for r in rows}
        legends = [lg for lg in legends if (lg.openalex_id or lg.name).lower() not in have]
        print(f" {len(legends)} new candidates. Fetching recent work...")
        L.add_recent_work(disc, legends)
        named_ids = {lg.openalex_id for lg in legends if lg.source == "you named them"}
        before = len(legends)
        # People you named are always kept; the rest must actually be physicists.
        legends = [lg for lg in legends if lg.openalex_id in named_ids or L.is_physics(lg)]
        print(f" Physics filter: kept {len(legends)} of {before}.")
        disc.save_cache()
        print(" Looking up public addresses (strict name checks)...")
        L.find_emails(legends, workers=args.workers)
        rows = L.merge(rows, legends)
        L.save(path, rows)
        ok = sum(1 for r in rows if r.get("Email"))
        print(f" {len(rows)} in {path.name}; {ok} with a verified address.")
        for r in rows[:25]:
            print(_safe(f"   h={str(r.get('h-index')):>3}  {str(r.get('Name'))[:28]:<28} "
                        f"{str(r.get('Institution'))[:34]:<34} {r.get('Email') or '-- no address'}"))

    if args.draft:
        _h("LEGACY RESEARCHERS: DRAFT")
        settings.require_api_key()
        ev = LLMEvaluator(settings, profile)
        n = 0
        for r in rows:
            if not r.get("Email") or r.get("Draft") or r.get("Send Status") == "sent":
                continue
            try:
                r["Subject"], r["Draft"] = L.draft(ev, settings, profile, r, args.summer)
                r["Why them"] = r.get("Source", "")
                n += 1
                print(_safe(f"\n  == {r['Name']} ({r['Institution']})\n  Subject: {r['Subject']}"))
                for line in r["Draft"].splitlines():
                    print(_safe("  | " + line))
            except Exception as exc:
                print(_safe(f"  ! {r['Name']}: {exc}"))
            L.save(path, rows)
        print(f"\n {n} letter(s) drafted. {ev.usage_report()}")
        print(f" Read them in {path}, set Approved = yes for the ones you want, then:")
        print("   python -m outreach.cli legacy --send            # dry run")
        print("   python -m outreach.cli legacy --send --live     # send from your Gmail")

    if args.send:
        _h(f"LEGACY RESEARCHERS: SEND ({'LIVE' if args.live else 'DRY RUN'})")
        s = settings
        alt = replace(s, smtp_host=s.resend_smtp_host, smtp_port=s.resend_smtp_port,
                      smtp_security=s.resend_smtp_security, smtp_user=s.resend_smtp_user,
                      smtp_password=s.resend_smtp_password, smtp_from=s.resend_smtp_from,
                      reply_to="")
        def approved(r):
            return str(r.get("Approved") or "").strip().lower() in {"y", "yes", "true", "1", "x"}
        # --approve-all only covers addresses the finder verified; anything typed in
        # by hand or unverified must be approved row by row.
        todo = [r for r in rows if r.get("Draft") and r.get("Email")
                and r.get("Send Status") != "sent" and "[FILL" not in r["Draft"]
                and (approved(r) or (args.approve_all and r.get("Email Status") == "ok"))]
        held = [r for r in rows if r.get("Draft") and r.get("Email") and r not in todo
                and r.get("Send Status") != "sent"]
        for r in held:
            print(_safe(f"  held: {r['Name']} <{r['Email']}> (status {r.get('Email Status')}; "
                        "set Approved = yes after checking)"))
        todo = todo[:args.limit]
        if not todo:
            print(" Nothing approved. Set Approved = yes in legacy.xlsx (or pass --approve-all).")
            return 0
        print(f" From {alt.smtp_from} to {len(todo)} researcher(s).")
        if args.live and not args.yes and input(" Send now? [Y/N]: ").strip().lower() not in {"y", "yes"}:
            print(" Aborted.")
            return 2
        import random, time
        with Mailer(alt, profile, dry_run=not args.live) as mailer:
            for i, r in enumerate(todo):
                msg = mailer.build_message(r["Email"], r["Name"], r["Subject"], r["Draft"])
                try:
                    mailer.deliver(msg)
                    if args.live:
                        r["Send Status"], r["Sent At"] = "sent", f"{datetime.now():%Y-%m-%d %H:%M:%S}"
                        r["Message ID"] = msg["Message-ID"]
                        append_send_log(settings.output_dir / "send_log.csv", "legacy",
                                        [{"sheet": "legacy", "row": i, "email": r["Email"],
                                          "name": r["Name"]}],
                                        [{"sheet": "legacy", "row": i, "status": "sent",
                                          "sent_at": r["Sent At"]}])
                    status = "sent" if args.live else "dry-run"
                except Exception as exc:
                    status = f"failed: {type(exc).__name__}"
                    r["Send Status"] = status
                print(_safe(f"  {status:<10} {r['Name'][:30]:<30} {r['Email']}"))
                L.save(path, rows)
                if args.live and i < len(todo) - 1:
                    Mailer._countdown(random.uniform(120, 300), "next letter")
    return 0


def cmd_replies(args) -> int:
    from .replier import ThreadReader, draft_reply, load_queue, plan_meeting, save_queue
    settings.ensure_dirs()
    queue_path = settings.output_dir / "reply_queue.json"
    profile = parse_resume(settings.require_resume(args.resume))

    if args.approve_all:
        existing = load_queue(queue_path)
        newly = 0
        for q in existing:
            if not q.get("approved") and not q.get("sent_at") and q.get("body") \
                    and q.get("action") != "none":
                q["approved"] = True
                newly += 1
        if newly:
            save_queue(queue_path, existing)
            print(f" Approved {newly} queued draft(s) in {queue_path.name}.")

    if args.send:
        items = [q for q in load_queue(queue_path)
                 if q.get("approved") and not q.get("sent_at") and q.get("body")]
        _h(f"SEND FOLLOW-UPS ({'LIVE' if args.live else 'DRY RUN'})")
        if not items:
            print(f" Nothing approved in {queue_path.name}. Set \"approved\": true on the"
                  " items you want, or pass --approve-all.")
            return 0
        window = None
        window_spec = (args.local_window if args.local_window is not None
                       else settings.send_local_window)
        if window_spec and window_spec.lower() not in {"off", "none", "any"}:
            try:
                window = SendWindow(window_spec, weekends=args.weekends)
            except ValueError as exc:
                print(f" ERROR: {exc}")
                return 1
            countries = _reply_countries(args)
            for q in items:
                if not q.get("country"):
                    q["country"] = countries.get(q["email"].lower()) or _tld_country(q["email"])
            print(f" Local window: {window.spec} recipient time"
                  f"{'' if window.weekends else ', weekdays only'}"
                  f"{' - waiting for it' if args.wait else ' - others held for a later run'}")
        if args.live and not args.yes and input(
                f" Send {len(items)} replies now? [Y/N]: ").strip().lower() not in {"y", "yes"}:
            print(" Aborted.")
            return 2
        all_items = load_queue(queue_path)
        held = []
        with Mailer(settings, profile, dry_run=not args.live) as mailer:
            pending = list(items)
            while pending:
                if window is not None:
                    pick = next((k for k, q in enumerate(pending) if window.ready(q)), None)
                    if pick is None:
                        waits = [window.seconds_until(q) for q in pending]
                        if args.wait and args.live:
                            mailer.close()
                            Mailer._countdown(min(waits), label="next local send window opens")
                            continue
                        held = list(zip(pending, waits))
                        break
                    q = pending.pop(pick)
                else:
                    q = pending.pop(0)
                err = ""
                try:
                    msg = mailer.build_reply(q["email"], q.get("name", ""), q["subject"],
                                             q["body"], q.get("in_reply_to", ""),
                                             q.get("references", ""))
                    mailer.deliver(msg)
                    status = "sent" if args.live else "dry-run"
                except Exception as exc:
                    status = f"failed: {type(exc).__name__}"
                    msg, err = None, str(exc)
                files = [p.get_filename() for p in msg.iter_attachments()] if msg else []
                extra = f"  [attached: {', '.join(files)}]" if files else ""
                print(_safe(f"  {status:<10} {q['email']}{extra}"))
                if err:
                    print(_safe(f"             -> {err[:160]}"))
                if status == "sent":
                    holds = _calendar_holds(q)
                    for a in all_items:
                        if _qkey(a) == _qkey(q):
                            a["sent_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                            if holds:
                                a["calendar_events"] = holds
                    # Persist each delivery so an interrupted --wait run can't resend.
                    save_queue(queue_path, all_items)
        for q, w in held[:5]:
            print(_safe(f"  held       {q['email']} ({q.get('country') or '?'}: "
                        f"opens in {fmt_duration(w)})"))
        if held:
            print(f"\n {len(held)} reply(s) held: outside their local {window.spec} window. "
                  "Re-run later, or add --wait to sit until each window opens.")
        save_queue(queue_path, all_items)
        return 0

    path = settings.require_target_xlsx(args.xlsx)
    rows = [r for r in load_for_sending(path, sheet=args.sheet)
            if r["send_status"].lower() == "sent"]
    _h("READ CONVERSATIONS + DRAFT REPLIES")
    if not rows:
        print(f" No rows marked 'sent' in {path.name}.")
        return 0

    def progress(folder, n, total, row):
        sys.stdout.write(f"\r  scanning {folder}: {n}/{total}   ")
        sys.stdout.flush()

    try:
        with ThreadReader(settings) as reader:
            threads = reader.threads(rows, progress=progress)
    except Exception as exc:
        print(f"\n ERROR: IMAP failed: {type(exc).__name__}: {exc}")
        return 1
    print()

    old = {_qkey(q): q for q in load_queue(queue_path)}
    _merge_sent_queue(threads, old.values())
    pending = [t for t in threads if t.awaiting_us()]
    print(f" {len(threads)} conversations | {len(pending)} where the professor wrote last")
    evaluator = LLMEvaluator(settings, profile)
    gcal_svc, taken = _gcal_or_none(), _held_slots(old.values())
    queue = list(old.values())
    for t in pending:
        last = t.last_incoming
        key = (t.row["email"], last.message_id, "reply", 0)
        if key in old and (not args.redraft or old[key].get("sent_at")):
            # Never redraft a reply that already went out: it would clear
            # sent_at and the next --send would mail the professor again.
            continue
        slots, slot_lines = plan_meeting(t, settings, gcal_svc, taken)
        try:
            d = draft_reply(evaluator, t, settings, profile, slot_lines)
        except Exception as exc:
            print(_safe(f"  ! {t.row['email']}: {exc}"))
            continue
        refs = " ".join(m.message_id for m in t.messages if m.message_id and m is not last)
        item = {"email": t.row["email"], "name": t.row.get("name", ""),
                "institution": t.row.get("institution", ""),
                "subject": last.subject or t.row.get("subject", ""),
                "in_reply_to": last.message_id, "references": refs,
                "their_message": last.text[:1500], "summary": d["summary"],
                "action": d["action"], "body": d["body"],
                "approved": bool(args.approve_all and d["action"] != "none"),
                "kind": "reply", "n": 0,
                "meet_slots": [s.isoformat() for s in slots] if d["action"] != "none" else [],
                "drafted_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "sent_at": ""}
        if item["meet_slots"]:
            taken += [(s, s + _meet_dur()) for s in slots]
        old[key] = item
        queue = list(old.values())
        save_queue(queue_path, queue)
        print(_safe(f"\n  [{d['action']}] {t.row.get('name')} <{t.row['email']}>"))
        print(_safe(f"    they said: {_safe(last.text[:160]).replace(chr(10), ' ')}"))
        print(_safe(f"    status   : {d['summary']}"))
        if d["body"]:
            for line in d["body"].splitlines():
                print(_safe("    | " + line))
    queue = _draft_funding_followups(threads, old, evaluator, profile, gcal_svc, taken,
                                     args, queue_path)
    save_queue(queue_path, queue)
    print(f"\n {evaluator.usage_report()}")
    print(f" Queue: {queue_path}")
    print(" Edit bodies / set \"approved\": true, then:")
    print("   python -m outreach.cli replies --send            # dry run")
    print("   python -m outreach.cli replies --send --live     # really send")
    return 0


def cmd_tls_check(args) -> int:
    import socket
    import ssl
    from urllib.parse import urlparse
    from .tls import context, describe
    _h("TLS CHECK")
    print(_safe(f" Python {sys.version.split()[0]} | {ssl.OPENSSL_VERSION}"))
    print(_safe(f" Trust source: {describe()}"))
    imap = settings.imap_host or settings.smtp_host.replace("smtp.", "imap.", 1)
    hosts = [(settings.smtp_host, settings.smtp_port), (imap, settings.imap_port),
             (urlparse(settings.base_url or "https://api.openai.com").hostname, 443),
             ("api.openalex.org", 443)]
    bad = 0
    for host, port in hosts:
        if not host:
            continue
        try:
            with socket.create_connection((host, port), timeout=15) as sock:
                if port in (587, 25):
                    print(_safe(f"   skip   {host}:{port} (STARTTLS port; tested at send)"))
                    continue
                with context().wrap_socket(sock, server_hostname=host) as tls:
                    print(_safe(f"   ok     {host}:{port}  {tls.version()}"))
        except Exception as exc:
            bad += 1
            print(_safe(f"   FAIL   {host}:{port}  {type(exc).__name__}: {exc}"))
    if bad:
        print("\n Fixes:  pip install -U truststore certifi")
        print("   macOS python.org build: run '/Applications/Python 3.x/Install Certificates.command'")
        print("   Behind a proxy/antivirus that inspects TLS: export SSL_CERT_FILE=/path/to/its-root.pem")
    return 1 if bad else 0


def cmd_check_replies(args) -> int:
    path = settings.require_target_xlsx(args.xlsx)
    rows = [r for r in load_for_sending(path, sheet=args.sheet)
            if r["send_status"].lower() == "sent"]
    _h("CHECK REPLIES")
    if not rows:
        print(f" No rows marked 'sent' in {path.name}.")
        return 0
    checker = ReplyChecker(settings)
    print(f" {len(rows)} sent emails | mailbox {checker.user} @ {checker.host}")

    def progress(folder, n, total, row):
        sys.stdout.write(f"\r  scanning {folder}: {n}/{total}   ")
        sys.stdout.flush()

    try:
        with checker:
            print(f" Folders: {', '.join(checker.folders)}")
            found = checker.check(rows, progress=progress)
    except Exception as exc:
        print(f"\n ERROR: IMAP failed: {type(exc).__name__}: {exc}")
        print(" Set IMAP_HOST / IMAP_USER / IMAP_PASSWORD in .env if they differ from SMTP.")
        return 1
    print()

    by_key = {(r["sheet"], r["row"]): r for r in rows}
    counts = Counter(rep.kind for rep in found.values())
    for kind in ("replied", "auto-reply", "bounced"):
        hits = [(by_key[k], rep) for k, rep in found.items() if rep.kind == kind]
        if not hits:
            continue
        print(f"\n {kind.upper()} ({len(hits)}):")
        for row, rep in sorted(hits, key=lambda h: h[1].when or datetime.min):
            when = f"{rep.when:%d %b %H:%M}" if rep.when else ""
            print(_safe(f"   {when:<13} {str(row['name'])[:28]:<28} {row['email'][:36]}"))
            if rep.snippet and kind == "replied":
                print(_safe(f"                 \"{rep.snippet[:110]}\""))
    waiting = len(rows) - len(found)
    print(f"\n Summary: replied={counts['replied']} auto-reply={counts['auto-reply']} "
          f"bounced={counts['bounced']} no answer yet={waiting}")

    changed = {k: rep for k, rep in found.items()
               if by_key[k].get("reply_status", "") != rep.status()}
    if args.dry_run or not changed:
        print(" Workbook unchanged." if not changed else " Dry run: workbook unchanged.")
        return 0
    if not is_writable(path):
        print(_safe(f" {path.name} is open in Excel; close it and re-run to save statuses."))
        return 1
    mark_replies(path, changed)
    print(f" Wrote {len(changed)} Reply Status update(s) to {path.name}.")
    return 0


def _verify_rows(professors, args_workers=None, repair=True, apply=True,
                 dns=True, probe=False, label=""):
    """Shared by verify/discover/recheck: parallel email check with progress."""
    cache_path = settings.output_dir / "email_lookup_cache.json"
    cache = load_cache(cache_path) if repair else {}
    workers = args_workers or settings.verify_workers
    done = [0]
    total = len(professors)

    def on_result(i, prof, res):
        done[0] += 1
        flag = {"ok": "ok", "suspect": "SUSPECT", "invalid": "INVALID"}.get(
            res.status, "?")
        print(_safe(f"  {label}[{done[0]}/{total}] {flag:<8} {res.confidence:>3}%  "
                    f"{prof.name[:28]:<28} {prof.email[:38]}"))
        if res.status != "ok":
            print(_safe(f"            {res.summary()[:100]}"))
        if res.suggestion and not apply:
            print(_safe(f"            -> suggested: {res.suggestion}  [{res.source}]"))
        if repair and done[0] % 25 == 0:
            save_cache(cache_path, cache)   # survive an interrupted run

    counts = verify_many(professors, workers=workers, check_dns=dns, probe=probe,
                         mail_from=settings.smtp_from, repair=repair, cache=cache,
                         apply=apply, on_result=on_result)
    if repair:
        save_cache(cache_path, cache)
    return counts


_DONE_STATUSES = {"sent", "review-sent"}


def _drop_sent(path: Path, professors):
    """Leave rows that were already mailed exactly as they are."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    done = set()
    for ws in wb.worksheets:
        hdr = _find_header_row(ws)
        col = _column_index(ws, hdr).get("send status")
        if not col:
            continue
        for r, row in enumerate(ws.iter_rows(min_row=hdr + 1, values_only=True), hdr + 1):
            val = row[col - 1] if col - 1 < len(row) else None
            if str(val or "").strip().lower() in _DONE_STATUSES:
                done.add((ws.title, r))
    wb.close()
    # Sends recorded while the workbook was locked live in the sidecar until
    # the next successful write; they are just as done.
    side = sidecar_path(path)
    if side.exists():
        try:
            for item in json.loads(side.read_text(encoding="utf-8")):
                if str(item.get("status", "")).strip().lower() in _DONE_STATUSES:
                    done.add((item.get("sheet"), item.get("row")))
        except (OSError, ValueError, AttributeError):
            pass
    keep = [p for p in professors if (p.source_sheet, p.source_row) not in done]
    skipped = len(professors) - len(keep)
    if skipped:
        print(f" Skipping {skipped} row(s) already sent - they are left untouched.")
    return keep


def cmd_discover(args) -> int:
    """Build a target list of N professors in K resumable batches."""
    settings.ensure_dirs()
    _h("DISCOVER TARGETS")

    top_n = (args.top_universities if args.top_universities is not None
             else settings.discover_top_universities)
    prestige = (args.prestige_weight if args.prestige_weight is not None
                else settings.discover_prestige_weight)
    if args.list_universities:
        from .discover import Discoverer, resolve_countries
        disc = Discoverer(mailto=settings.openalex_mailto, api_key=settings.openalex_api_key,
                          cache_path=settings.output_dir / "discover_cache.json")
        countries = args.countries or settings.discover_countries
        unis = disc.top_universities(resolve_countries(countries), top_n or 30,
                                     args.since or settings.discover_since)
        disc.save_cache()
        if not unis:
            print(f" No universities found. {disc.last_error}")
            return 1
        print(f" Top {len(unis)} universities in {countries} "
              "(by highly-cited papers since " f"{args.since or settings.discover_since}):")
        for i, u in enumerate(unis, 1):
            print(_safe(f"  {i:>3}. {u['name']:<55} {u['country']}  {u['papers']:>6}"))
        return 0

    queries = split_queries(args.query or settings.discover_queries)
    if not queries:
        # Default to what the resume says you actually work on; that is the
        # whole point of having parsed it. One interest per batch.
        profile = parse_resume(settings.require_resume(args.resume))
        queries = [q for q in profile.research_interests if len(q) > 3][:8] \
            or [" ".join(profile.keywords[:8])]
        print(_safe(" No --query given; using your resume interests:\n   "
                    + "; ".join(queries)))
    if not any(q.strip() for q in queries):
        print(" Nothing to search for. Pass --query \"topic one; topic two\".")
        return 1

    count = args.count or settings.discover_count
    n_batches = args.batches or settings.discover_batches
    countries = args.countries or settings.discover_countries
    since = args.since or settings.discover_since
    max_pages = args.max_pages or settings.discover_max_pages

    out_path = Path(args.output) if args.output else settings.output_dir / "professors.xlsx"
    state_path = out_path.with_suffix(".batches.json")
    if args.fresh:
        state_path.unlink(missing_ok=True)
    existing = load_targets(out_path) if out_path.exists() else []
    excluded = []
    for ex in args.exclude or []:
        if Path(ex).exists():
            excluded += load_targets(ex)
        else:
            print(_safe(f" WARNING: --exclude {ex} not found; nobody excluded from it."))
    if excluded:
        print(f" Excluding {len(excluded)} professors listed in --exclude workbooks.")
    if existing and not is_writable(out_path):
        print(_safe(f" ERROR: {out_path.name} is open in Excel. Close it and re-run."))
        return 1

    plan = plan_batches(queries, count, n_batches)
    print(f" Target: {count} professors in {n_batches} batch(es) -> {out_path}")
    print(f" Region: {countries}   papers since {since}   seniority: "
          f"h-index>={args.min_h_index}, works>={args.min_works}")
    if top_n:
        print(f" Universities: only the region's top {top_n}, prestige weight {prestige:g}")
    for spec in plan:
        print(_safe(f"   batch {spec.index}: {spec.quota:>4} x '{spec.query}'"))
    if existing:
        print(f" {len(existing)} professors already in {out_path.name}; "
              "they will not be added again.")

    def on_start(spec):
        _h(f"BATCH {spec.index}/{n_batches}: {spec.query}  (quota {spec.quota})")

    def on_page(spec, page, n, total, found):
        print(f"  page {page:>2}  +{n} papers  (topic has {total:,})  "
              f"senior researchers so far: {found}/{spec.quota}")

    def on_top_unis(unis, err):
        if not unis:
            print(f" ! Could not rank universities: {err or 'no results'}")
            return
        print(_safe(" Top universities: " + ", ".join(u["name"] for u in unis[:12])
                    + (f" ... (+{len(unis) - 12})" if len(unis) > 12 else "")))

    totals = {"found": 0}
    rows_now = [len(existing)]
    for profs, res in run_batches(
            queries, count=count, batches=n_batches, countries=countries,
            since_year=since, min_h_index=args.min_h_index,
            min_works=args.min_works, max_pages_per_batch=max_pages,
            existing=existing + excluded, state_path=state_path,
            mailto=settings.openalex_mailto, api_key=settings.openalex_api_key,
            cache_path=settings.output_dir / "discover_cache.json",
            top_universities=top_n, prestige_weight=prestige,
            also_universities=split_queries(args.also_universities
                                            or settings.discover_also_universities),
            on_top_unis=on_top_unis,
            on_batch_start=on_start, on_page=on_page):
        if res.error:
            print(f" ! {res.error}")
        why = ("topic exhausted" if res.exhausted
               else "page cap reached" if res.found < res.quota else "quota met")
        print(f" Batch {res.index}: {res.found}/{res.quota} new professors "
              f"from {res.works} papers ({why})")
        if not profs:
            continue
        if args.verify_emails:
            print(f" Finding current email addresses for {len(profs)} people...")
            c = _verify_rows(profs, args.workers, repair=True, apply=True,
                             label=f"b{res.index} ")
            print(f" Emails: ok={c.get('ok', 0)} suspect={c.get('suspect', 0)} "
                  f"not found={c.get('invalid', 0)}")
        if not _save_batch(out_path, profs):
            return 1
        totals["found"] += len(profs)
        rows_now[0] = len(load_targets(out_path))
        print(_safe(f" Saved -> {out_path.name} ({rows_now[0]} rows total)"))

    final = rows_now[0]
    _h("DISCOVERY DONE")
    print(f" {totals['found']} new professors this run; {final} in {out_path.name}.")
    if final < count:
        print(f" Short of {count}: widen --countries, lower --min-h-index, use an "
              "earlier --since, or add more topics to --query.")
    print("\n Next:")
    if not args.verify_emails:
        print(f"   python -m outreach.cli verify --xlsx \"{out_path}\" --repair --apply")
    print(f"   python -m outreach.cli run --mode list --xlsx \"{out_path}\"")
    return 0


def _save_batch(out_path: Path, profs) -> bool:
    """Append one batch, waiting out an Excel lock instead of losing it.

    The batch is only marked done after this returns, so if the user gives up
    the rows are also parked in a JSON file and the batch reruns next time.
    """
    import time
    while out_path.exists() and not is_writable(out_path):
        print(_safe(f" {out_path.name} is open in Excel - close it to save this "
                    "batch (retrying in 15s, Ctrl+C to stop)."))
        try:
            time.sleep(15)
        except KeyboardInterrupt:
            break
    try:
        append_rows(out_path, profs)
        return True
    except (PermissionError, OSError) as exc:
        dump = out_path.with_suffix(".unsaved.json")
        dump.write_text(json.dumps([p.to_dict() if hasattr(p, "to_dict") else vars(p)
                                    for p in profs], indent=1, default=str),
                        encoding="utf-8")
        print(_safe(f" ERROR: could not write {out_path.name} ({exc}). "
                    f"This batch was parked in {dump.name} and will be re-run "
                    "next time - re-run the same command once the file is closed."))
        return False


def cmd_recheck(args) -> int:
    """Re-verify every professor (still active, still there) and every email."""
    settings.ensure_dirs()
    path = settings.require_target_xlsx(args.xlsx)
    _h("RECHECK PROFESSORS AND EMAILS")
    professors = load_targets(path, sheet=args.sheet)
    print(f" {len(professors)} rows in {path.name}")
    if not args.include_sent:
        professors = _drop_sent(path, professors)

    if not args.emails_only:
        print("\n Professors: still publishing, still at the listed institution?")
        client = OpenAlexClient(mailto=settings.openalex_mailto,
                                api_key=settings.openalex_api_key,
                                cache_path=settings.output_dir / "openalex_cache.json")

        def on_prof(i, prof, status, note):
            print(_safe(f"  [{i + 1}/{len(professors)}] {status or '?':<9} "
                        f"{prof.name[:28]:<28} {note[:70]}"))

        pc = recheck_all(professors, client, on_result=on_prof)
        print(f" active={pc.get('active', 0)} moved={pc.get('moved', 0)} "
              f"inactive={pc.get('inactive', 0)} not-found={pc.get('not found', 0)}")

    if not args.professors_only:
        print("\n Emails: looking up each person's current published address")
        ec = _verify_rows(professors, args.workers, repair=True, apply=args.apply,
                          dns=not args.no_dns)
        print(f" ok={ec.get('ok', 0)} suspect={ec.get('suspect', 0)} "
              f"invalid={ec.get('invalid', 0)}")

    if args.dry_run:
        print(" Dry run: nothing written.")
        return 0
    if not is_writable(path):
        print(_safe(f" ERROR: {path.name} is open in Excel. Close it and re-run."))
        return 1
    if args.apply:
        _write_emails(path, professors)
    update_workbook(path, professors, make_backup=not args.no_backup)
    print(f" Wrote Profile Status and Email Status columns -> {path.name}")
    if not args.apply:
        print(" Re-run with --apply to replace addresses with the suggested ones.")
    return 0


def _write_emails(path: Path, professors) -> None:
    """Write adopted addresses into the Email column itself."""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    for p in professors:
        if not (p.source_sheet and p.source_row and p.email):
            continue
        ws = wb[p.source_sheet]
        hdr = _find_header_row(ws)
        cols = _column_index(ws, hdr)
        col = next((cols[h] for h in ("email address", "email", "contact email")
                    if h in cols), None)
        if col:
            ws.cell(row=p.source_row, column=col, value=p.email)
    wb.save(path)


# ---------------------------------------------------------------- parser
def cmd_verify(args) -> int:
    """Check every address, and optionally hunt down the correct one."""
    settings.ensure_dirs()
    path = settings.require_target_xlsx(args.xlsx)
    _h("EMAIL VERIFICATION")
    professors = load_targets(path, sheet=args.sheet)
    if not args.include_sent:
        professors = _drop_sent(path, professors)
    if args.only_missing:
        # Rows that already have a checked address don't need another web search.
        before = len(professors)
        professors = [p for p in professors if not (p.email or "").strip()]
        print(f" --only-missing: {before - len(professors)} rows already have an "
              "address; looking up the rest.")
        # Never-searched rows (e.g. fresh `discover` output) first; rows that
        # already failed a search are the least likely to succeed now.
        from openpyxl import load_workbook
        tried = set()
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            hdr = _find_header_row(ws)
            col = _column_index(ws, hdr).get("email status")
            if not col:
                continue
            for r, row in enumerate(ws.iter_rows(min_row=hdr + 1, values_only=True), hdr + 1):
                if col - 1 < len(row) and row[col - 1]:
                    tried.add((ws.title, r))
        wb.close()
        professors.sort(key=lambda p: (p.source_sheet, p.source_row) in tried)
        fresh = sum(1 for p in professors if (p.source_sheet, p.source_row) not in tried)
        print(f" {fresh} never searched (done first), {len(professors) - fresh} retried after.")
    if args.top:
        # Rank like `send` does: by the tool's Probability of Acceptance,
        # falling back to a hand-typed Probability column.
        scored = {(r["sheet"], r["row"]): r["probability"]
                  for r in load_for_sending(path, sheet=args.sheet)}
        professors = sorted(
            professors,
            key=lambda p: -max(_pct(p.prior_probability),
                               scored.get((p.source_sheet, p.source_row), 0)))[:args.top]
    print(f" Checking {len(professors)} addresses from {path.name}")
    print(" Checks: syntax, mail records (MX/A), name match")
    if args.repair:
        print(" Repair on: reading each professor's own pages for their latest address")
    if args.probe_smtp:
        print(" SMTP probe on: advisory only, many servers answer for every mailbox")

    counts = _verify_rows(professors, args.workers, repair=args.repair,
                          apply=args.apply, dns=not args.no_dns,
                          probe=args.probe_smtp)
    print(f"\n ok={counts.get('ok', 0)} suspect={counts.get('suspect', 0)} "
          f"invalid={counts.get('invalid', 0)}")
    if args.dry_run:
        print(" Dry run: nothing written to the workbook.")
        return 0
    if not is_writable(path):
        print(_safe(f" ERROR: {path.name} is open in Excel. Close it and re-run."))
        return 1
    if args.apply:
        _write_emails(path, professors)
    update_workbook(path, professors, make_backup=not args.no_backup)
    print(f" Wrote Email Status / Confidence / Notes / Suggested Email -> {path.name}")
    if counts.get("invalid"):
        print(" Rows marked invalid are skipped by `send` unless you pass "
              "--allow-unverified.")
    if not args.apply and any(p.email_suggestion for p in professors):
        print(" Review the Suggested Email column, then re-run with --apply to "
              "adopt the suggestions.")
    return 0


def _pct(value) -> int:
    try:
        return int(float(str(value).strip().rstrip("%")))
    except (TypeError, ValueError):
        return 0


def cmd_models(args) -> int:
    if args.provider:
        settings.provider = args.provider
    _h(f"AVAILABLE MODELS ({settings.provider})")
    try:
        models = list_models(settings)
    except ConfigError as exc:
        print(f" {exc}")
        return 1
    chat = [m for m in models if not any(
        t in m for t in ("whisper", "orpheus", "guard", "tts", "embedding"))]
    for m in chat:
        marker = " <- current DRAFT_MODEL" if m == settings.draft_model else ""
        price = "free" if settings.is_free_tier else (
            f"${MODEL_PRICES[m][0]:.2f}/{MODEL_PRICES[m][2]:.2f} per 1M in/out"
            if m in MODEL_PRICES else "pricing unknown")
        print(_safe(f"  {m:<42} {price}{marker}"))
    other = [m for m in models if m not in chat]
    if other:
        print(f"\n Non-chat models (speech/moderation), not usable here: "
              f"{', '.join(other)}")
    print("\n Set DRAFT_MODEL in .env, or pass --model to `run`.")
    return 0


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="outreach",
        description="Resume-driven professor matching and personalized cold-email automation.")
    sub = p.add_subparsers(dest="command", required=True)

    sp = sub.add_parser("profile", help="Parse the resume and show what was extracted")
    sp.add_argument("--resume", help="Path to resume PDF/DOCX")
    sp.set_defaults(func=cmd_profile)

    sm = sub.add_parser("models", help="List the models your provider currently offers")
    sm.add_argument("--provider", choices=sorted(PROVIDERS),
                    help="Provider to query (default: LLM_PROVIDER from .env)")
    sm.set_defaults(func=cmd_models)

    sr = sub.add_parser("run", help="Run the full match + draft pipeline")
    sr.add_argument("--mode", choices=["list", "scrape"], default="list",
                    help="list = Mode B (existing xlsx); scrape = Mode A (discovery)")
    sr.add_argument("--resume", help="Path to resume PDF/DOCX")
    sr.add_argument("--xlsx", help="Target workbook (Mode B)")
    sr.add_argument("--sheet", help="Restrict to a single sheet name")
    sr.add_argument("--sources", default="sources.yaml",
                    help="YAML/JSON list of directory URLs (Mode A)")
    sr.add_argument("--output", help="Write results to a new workbook instead of in place")
    sr.add_argument("--max-professors", type=int, help="Cap on LLM-evaluated professors")
    sr.add_argument("--min-score", type=float, help="Local match score threshold (0-1)")
    sr.add_argument("--budget", type=float, help="Hard USD budget for this run")
    sr.add_argument("--provider", choices=sorted(PROVIDERS),
                    help="LLM provider (default: LLM_PROVIDER from .env)")
    sr.add_argument("--model", help="Override the drafting model")
    sr.add_argument("--no-enrich", action="store_true", help="Skip OpenAlex lookups")
    sr.add_argument("--skip", action="append", default=[], metavar="STEP",
                    help="Skip optional steps (comma-separated or repeated): "
                         f"{', '.join(SKIPPABLE)}. e.g. --skip enrich,filter")
    sr.add_argument("--no-profile-hop", action="store_true",
                    help="Mode A: do not visit individual profile pages for emails")
    sr.add_argument("--no-backup", action="store_true", help="Do not back up the workbook")
    sr.add_argument("--no-focus", action="store_true",
                    help="Ignore FOCUS and include rows outside your field")
    sr.add_argument("--yes", action="store_true", help="Skip the [Y/N] cost gate")
    sr.set_defaults(func=cmd_run)

    ss = sub.add_parser("send", help="Dispatch the drafted emails")
    ss.add_argument("--xlsx", help="Workbook containing the drafts")
    ss.add_argument("--sheet", help="Restrict to a single sheet name")
    ss.add_argument("--resume", help="Path to resume (for the signature block)")
    ss.add_argument("--live", action="store_true", help="Actually send (default is dry-run)")
    ss.add_argument("--dry-run", action="store_true", help="Explicit dry run (default)")
    ss.add_argument("--review", action="store_true",
                    help="Mail every draft to yourself (SMTP_FROM) instead of professors")
    ss.add_argument("--review-to", help="Mail every draft to this address instead")
    ss.add_argument("--export", action="store_true",
                    help="Write each draft as a .eml file you can open in a mail client")
    ss.add_argument("--export-dir", help="Where to write .eml files (default output/drafts)")
    ss.add_argument("--resend", action="store_true", help="Re-send rows already marked sent")
    ss.add_argument("--top", type=int,
                    help="Mail only the N highest-probability professors")
    ss.add_argument("--limit", type=int, help="Max emails this run")
    ss.add_argument("--delay-min", type=float, help="Min seconds between emails")
    ss.add_argument("--delay-max", type=float, help="Max seconds between emails")
    ss.add_argument("--min-probability", type=int,
                    help="Only send to rows at or above this probability")
    ss.add_argument("--preview", type=int, default=2,
                    help="How many full drafts to print in dry-run")
    ss.add_argument("--allow-unverified", action="store_true",
                    help="Send even to rows `verify` marked invalid")
    ss.add_argument("--at", metavar="TIME",
                    help="Wait, then start sending at HH:MM (or 'YYYY-MM-DD HH:MM'), "
                         "your local time")
    ss.add_argument("--local-window", metavar="H-H",
                    help="Only send inside the recipient's local hours, e.g. 9-11 "
                         "(Country column; default SEND_LOCAL_WINDOW; 'off' disables)")
    ss.add_argument("--wait", action="store_true",
                    help="With a local window, keep running until every row's window "
                         "opens instead of holding them for a later run")
    ss.add_argument("--weekends", action="store_true",
                    help="Allow the local window on Saturday/Sunday")
    ss.add_argument("--no-focus", action="store_true",
                    help="Ignore FOCUS and include rows outside your field")
    ss.add_argument("--yes", action="store_true",
                    help="Skip the [Y/N] confirmation (for Task Scheduler / cron)")
    ss.set_defaults(func=cmd_send)

    st = sub.add_parser("stats", help="Count rows drafted, sent, replied, skipped (and why)")
    st.add_argument("--xlsx", help="Workbook to summarise")
    st.add_argument("--sheet", help="Restrict to a single sheet name")
    st.set_defaults(func=cmd_stats)

    tc = sub.add_parser("tls-check", help="Diagnose certificate errors (e.g. on macOS)")
    tc.set_defaults(func=cmd_tls_check)

    rp = sub.add_parser("replies",
                        help="Read each professor's full thread and draft follow-up replies")
    rp.add_argument("--xlsx", help="Workbook containing the sent rows")
    rp.add_argument("--sheet", help="Restrict to a single sheet name")
    rp.add_argument("--resume", help="Path to resume")
    rp.add_argument("--redraft", action="store_true", help="Redraft already-queued replies")
    rp.add_argument("--approve-all", action="store_true",
                    help="Approve every unsent reply/thank draft (queued and new)")
    rp.add_argument("--send", action="store_true", help="Send approved replies from the queue")
    rp.add_argument("--live", action="store_true", help="With --send: actually send")
    rp.add_argument("--yes", action="store_true", help="Skip confirmation")
    rp.add_argument("--local-window", metavar="H-H",
                    help="With --send: only send inside the recipient's local hours, e.g. 9-17 "
                         "(default SEND_LOCAL_WINDOW; 'off' disables)")
    rp.add_argument("--wait", action="store_true",
                    help="With --local-window: keep running until every window opens")
    rp.add_argument("--weekends", action="store_true",
                    help="Allow the local window on Saturday/Sunday")
    rp.set_defaults(func=cmd_replies)

    ar = sub.add_parser("archive-sent",
                        help="Copy every email sent over SMTP into your IMAP Sent folder")
    ar.add_argument("--xlsx", help="Workbook containing the sent rows")
    ar.add_argument("--resume", help="Path to resume")
    ar.add_argument("--live", action="store_true", help="Actually copy (default: dry run)")
    ar.set_defaults(func=cmd_archive_sent)

    cv = sub.add_parser("conversations",
                        help="Build output/conversations.html with every full thread")
    cv.add_argument("--xlsx", help="Workbook containing the sent rows")
    cv.add_argument("--no-open", action="store_true", help="Don't open it in the browser")
    cv.set_defaults(func=cmd_conversations)

    rf = sub.add_parser("replace-failed",
                        help="Mark bounced/wrong-person addresses invalid and add a colleague "
                             "from the same department for each")
    rf.add_argument("--xlsx", help="Workbook (default TARGET_XLSX)")
    rf.add_argument("--per-failure", type=int, default=1,
                    help="Colleagues to add per unreachable professor (default 1)")
    rf.set_defaults(func=cmd_replace_failed)

    lg = sub.add_parser("legacy", aliases=["legacy-researchers"],
                        help="Eminent researchers in your field + a humble aspiring-scientist "
                             "letter, sent from your personal Gmail (RESEND_SMTP_*)")
    lg.add_argument("--find", action="store_true", help="Search for candidates (default on first run)")
    lg.add_argument("--topics", help="';'-separated fields (default: particle physics, HEP pheno, "
                                     "ML for particle physics, cosmology, QFT, astrophysics)")
    lg.add_argument("--include", metavar="NAMES",
                    help="People you name yourself, e.g. \"Neil deGrasse Tyson; Fabiola Gianotti\"")
    lg.add_argument("--countries", help="Restrict region (default: worldwide)")
    lg.add_argument("--min-h", type=int, default=60, help="Minimum h-index (default 60)")
    lg.add_argument("--per-topic", type=int, default=8, help="Candidates per topic (default 8)")
    lg.add_argument("--workers", type=int, default=4, help="Parallel email lookups")
    lg.add_argument("--draft", action="store_true", help="Write a letter for each with an address")
    lg.add_argument("--summer", default="2027", help="Which summer (default 2027)")
    lg.add_argument("--send", action="store_true", help="Send approved letters")
    lg.add_argument("--approve-all", action="store_true", help="Treat every drafted letter as approved")
    lg.add_argument("--limit", type=int, default=10, help="Max letters per run (default 10)")
    lg.add_argument("--live", action="store_true", help="With --send: actually send")
    lg.add_argument("--yes", action="store_true", help="Skip confirmation")
    lg.add_argument("--resume", help="Path to resume")
    lg.set_defaults(func=cmd_legacy)

    ca = sub.add_parser("calendar-auth",
                        help="One-time Google Calendar sign-in for meeting holds")
    ca.set_defaults(func=cmd_calendar_auth)

    rs = sub.add_parser("resend",
                        help="Re-deliver the original email from a second account "
                             "(RESEND_SMTP_*) to bounced / unanswered professors")
    rs.add_argument("--who", choices=["all", "bounced", "noreply"], default="all",
                    help="bounced, noreply (sent, no answer), or all (default)")
    rs.add_argument("--xlsx", help="Workbook containing the sent rows")
    rs.add_argument("--sheet", help="Restrict to a single sheet name")
    rs.add_argument("--resume", help="Path to resume")
    rs.add_argument("--min-days", type=int,
                    help="Skip if the first email is newer than this (default RESEND_MIN_DAYS=4)")
    rs.add_argument("--limit", type=int, help="Max resends per day (default RESEND_DAILY_CAP)")
    rs.add_argument("--local-window", metavar="H-H",
                    help="Recipient local hours, e.g. 9-17 (default SEND_LOCAL_WINDOW)")
    rs.add_argument("--wait", action="store_true", help="Wait for each window to open")
    rs.add_argument("--weekends", action="store_true", help="Allow Saturday/Sunday")
    rs.add_argument("--prioritize-top", type=int, default=50, metavar="N",
                    help="Send to the region's N best universities first (default 50, 0 = off)")
    rs.add_argument("--region", help="Region for --prioritize-top (default DISCOVER_COUNTRIES)")
    rs.add_argument("--delay-min", type=float, help="Min seconds between emails (default SEND_DELAY_MIN)")
    rs.add_argument("--delay-max", type=float, help="Max seconds between emails (default SEND_DELAY_MAX)")
    rs.add_argument("--live", action="store_true", help="Actually send (default: dry run)")
    rs.add_argument("--yes", action="store_true", help="Skip confirmation")
    rs.set_defaults(func=cmd_resend)

    sy = sub.add_parser("check-replies",
                        help="Scan your mailbox (IMAP) for replies, auto-replies and bounces")
    sy.add_argument("--xlsx", help="Workbook containing the sent rows")
    sy.add_argument("--sheet", help="Restrict to a single sheet name")
    sy.add_argument("--dry-run", action="store_true", help="Report only")
    sy.set_defaults(func=cmd_check_replies)

    sd = sub.add_parser("discover",
                        help="Find N professors from research topics, in K batches")
    sd.add_argument("--query", help="Topics separated by ';' - one per batch, reused "
                                    "round-robin (default: DISCOVER_QUERIES, "
                                    "then your resume's interests)")
    sd.add_argument("--resume", help="Path to resume PDF/DOCX (for the default query)")
    sd.add_argument("--count", type=int,
                    help="Total professors to find (default DISCOVER_COUNT=1000)")
    sd.add_argument("--batches", type=int,
                    help="Split the search into this many batches (default 4)")
    sd.add_argument("--countries",
                    help="Region or ISO codes: "
                         + ", ".join(sorted(COUNTRY_GROUPS)) + ", or e.g. de,ch,gb")
    sd.add_argument("--since", type=int,
                    help="Only count papers published from this year on")
    sd.add_argument("--max-pages", type=int,
                    help="Cap on pages of 100 papers per batch (default 40)")
    sd.add_argument("--min-h-index", type=int, default=8,
                    help="Seniority gate: skip people below this h-index")
    sd.add_argument("--min-works", type=int, default=20,
                    help="Seniority gate: skip people with fewer publications")
    sd.add_argument("--top-universities", type=int, metavar="N",
                    help="Only look at the N best universities in --countries "
                         "(ranked by highly-cited output), and favour higher-ranked "
                         "ones (default DISCOVER_TOP_UNIVERSITIES, 0 = off)")
    sd.add_argument("--prestige-weight", type=float,
                    help="With --top-universities: how much university rank counts "
                         "against topic fit. 0 = fit only, 1 = balanced (default), "
                         "3 = mostly prestige")
    sd.add_argument("--also-universities", metavar="NAMES",
                    help="';'-separated universities to include too, e.g. "
                         "\"TU Berlin; EPFL\" (default DISCOVER_ALSO_UNIVERSITIES)")
    sd.add_argument("--list-universities", action="store_true",
                    help="Just print the top universities for --countries and exit")
    sd.add_argument("--output",
                    help="Workbook to grow (default output/professors.xlsx)")
    sd.add_argument("--verify-emails", action="store_true",
                    help="Find each person's current email as every batch completes")
    sd.add_argument("--workers", type=int, help="Parallel email lookups (default 4)")
    sd.add_argument("--exclude", action="append", metavar="XLSX",
                    help="Skip anyone already in this workbook (repeatable)")
    sd.add_argument("--fresh", action="store_true",
                    help="Forget saved batch progress and start the plan over")
    sd.set_defaults(func=cmd_discover)

    sc = sub.add_parser("recheck",
                        help="Re-verify every professor (active? moved?) and every email")
    sc.add_argument("--xlsx", help="Workbook to recheck (default TARGET_XLSX)")
    sc.add_argument("--sheet", help="Restrict to a single sheet name")
    sc.add_argument("--apply", action="store_true",
                    help="Replace addresses with the current ones found online")
    sc.add_argument("--professors-only", action="store_true",
                    help="Skip the email lookups")
    sc.add_argument("--emails-only", action="store_true",
                    help="Skip the OpenAlex professor check")
    sc.add_argument("--workers", type=int, help="Parallel email lookups (default 4)")
    sc.add_argument("--no-dns", action="store_true", help="Skip MX/A record lookups")
    sc.add_argument("--include-sent", action="store_true",
                    help="Also recheck rows already marked sent (skipped by default)")
    sc.add_argument("--dry-run", action="store_true", help="Report only")
    sc.add_argument("--no-backup", action="store_true", help="Do not back up the workbook")
    sc.set_defaults(func=cmd_recheck)

    sv = sub.add_parser("verify", help="Check (and optionally repair) email addresses")
    sv.add_argument("--xlsx", help="Workbook containing the targets")
    sv.add_argument("--sheet", help="Restrict to a single sheet name")
    sv.add_argument("--top", type=int, help="Only check the N highest-probability rows")
    sv.add_argument("--repair", action="store_true",
                    help="Search profile pages and the web for the correct address")
    sv.add_argument("--apply", action="store_true",
                    help="Overwrite the Email column with the suggested address")
    sv.add_argument("--probe-smtp", action="store_true",
                    help="Ask the mail server whether the mailbox exists (advisory)")
    sv.add_argument("--no-dns", action="store_true", help="Skip MX/A record lookups")
    sv.add_argument("--include-sent", action="store_true",
                    help="Also check rows already marked sent (skipped by default)")
    sv.add_argument("--workers", type=int, help="Parallel lookups (default VERIFY_WORKERS=4)")
    sv.add_argument("--only-missing", action="store_true",
                    help="Only look up rows with no email address yet")
    sv.add_argument("--dry-run", action="store_true",
                    help="Report only, do not touch the workbook")
    sv.add_argument("--no-backup", action="store_true", help="Do not back up the workbook")
    sv.set_defaults(func=cmd_verify)
    return p


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    try:
        return args.func(args)
    except KeyboardInterrupt:
        print("\n Interrupted.")
        return 130
    except ConfigError as exc:
        print(f"\n CONFIG: {exc}")
        return 1
    except Exception as exc:
        print(f"\n ERROR: {type(exc).__name__}: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
