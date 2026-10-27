"""Replace professors we could not reach with a colleague at the same institution.

For every sent row that hard-bounced or reached the wrong person:
  1. mark the address invalid in the workbook (so `send`/`resend` skip it), and
  2. find ONE senior researcher at the same institution who publishes on the
     same topics (via OpenAlex), and append them as a new row to be verified,
     drafted and sent like any other.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

from .batches import name_key
from .discover import Candidate, Discoverer, to_professor
from .excel_io import _column_index, _find_header_row, append_rows, load_targets

HARD = {"WRONG_PERSON_NAME_COLLISION", "MAILBOX_DOES_NOT_EXIST", "DEAD_DOMAIN",
        "RECIPIENT_REJECTED_BY_POLICY", "BOUNCED_OTHER"}


def mark_invalid(path: Path, failures: list[dict]) -> int:
    by_email = {f["email"].lower(): f for f in failures if f["root_cause"] in HARD}
    wb = load_workbook(path)
    n = 0
    for ws in wb.worksheets:
        hdr = _find_header_row(ws)
        cols = _column_index(ws, hdr)
        if "email address" not in cols or "email status" not in cols:
            continue
        for r in range(hdr + 1, ws.max_row + 1):
            e = str(ws.cell(r, cols["email address"]).value or "").strip().lower()
            f = by_email.get(e)
            if not f:
                continue
            ws.cell(r, cols["email status"]).value = "invalid"
            if "email confidence" in cols:
                ws.cell(r, cols["email confidence"]).value = "0%"
            if "email notes" in cols:
                ws.cell(r, cols["email notes"]).value = (
                    f"{f['root_cause']}: {f['server_or_professor_said'][:180]}")
            n += 1
    wb.save(path)
    return n


def _institution_id(disc: Discoverer, prof) -> str:
    if prof.openalex_id:
        aid = prof.openalex_id.rsplit("/", 1)[-1]
        data = disc._api(f"/authors/{aid}", {"select": "last_known_institutions"}) or {}
        insts = data.get("last_known_institutions") or []
        for i in insts:
            if (i.get("display_name") or "").lower() == (prof.institution or "").lower():
                return i["id"]
        if insts:
            return insts[0]["id"]
    hit = disc.find_universities([prof.institution]) if prof.institution else []
    return hit[0]["id"] if hit else ""


def _topic_query(prof) -> str:
    topics = [t.strip() for t in re.split(r"[;,]", prof.research_area or "") if t.strip()]
    return " ".join(topics[:2]) or (prof.recent_works[0] if prof.recent_works else "")


def find_colleague(disc: Discoverer, prof, skip_ids: set[str], skip_names: set[str],
                   min_h: int = 12, since: int = 2021):
    inst = _institution_id(disc, prof)
    query = _topic_query(prof)
    if not inst or not query:
        return None, "no institution id or topic"
    people: dict[str, Candidate] = {}
    for page in range(3):
        works, cursor, _ = disc.fetch_page(query, [], since, inst_ids=[inst],
                                           cursor="*" if page == 0 else cursor)
        if not works:
            break
        ranks = {inst: 1}
        for w in works:
            disc.fold(w, people, [], only_insts=ranks)
        if not cursor:
            break
    want_inst = _norm(prof.institution)
    want_topics = _words(" ".join([prof.research_area or "", query]))
    cands = [c for aid, c in people.items()
             if aid not in skip_ids and name_key(c.name) not in skip_names
             and _real_name(c.name)
             # Same institution only: a "colleague" elsewhere isn't a colleague.
             and _norm(c.institution) == want_inst]
    disc.hydrate(cands)
    good = []
    for c in cands:
        if c.h_index < min_h or c.works_count < 30 or c.hits < 2:
            continue
        overlap = len(want_topics & _words(" ".join(c.topics + c.titles)))
        if overlap >= 3:           # genuinely works on the same things
            good.append((overlap, c))
    if not good:
        return None, f"no same-department colleague on '{query[:40]}'"
    best = max(good, key=lambda oc: (oc[0], oc[1].score))[1]
    return best, query


_STOP = {"and", "of", "the", "in", "for", "with", "on", "a", "to", "using", "learning",
         "analysis", "applications", "methods", "systems", "techniques", "data", "based"}


def _words(text: str) -> set[str]:
    return {w for w in re.findall(r"[a-z]{4,}", (text or "").lower()) if w not in _STOP}


def _norm(name: str) -> str:
    import unicodedata
    s = unicodedata.normalize("NFKD", name or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z]", "", s)


def _real_name(name: str) -> bool:
    bad = ("proceedings", "consortium", "collaboration", "group", "team", "university")
    return bool(name) and not any(b in name.lower() for b in bad) and len(name.split()) >= 2


def replace_failed(settings, path: Path, failures_path: Path, limit: int = 1,
                   progress=print) -> dict:
    data = json.loads(failures_path.read_text(encoding="utf-8"))
    failures = [f for f in data["failures"] if f["root_cause"] in HARD]
    marked = mark_invalid(path, failures)
    progress(f" Marked {marked} unreachable address(es) invalid in {path.name}.")

    profs = load_targets(path)
    by_email = {(p.email or "").lower(): p for p in profs}
    skip_ids = {p.openalex_id for p in profs if p.openalex_id}
    skip_names = {name_key(p.name) for p in profs}
    already = {str(p.notes or "") for p in profs}
    disc = Discoverer(mailto=settings.openalex_mailto, api_key=settings.openalex_api_key,
                      cache_path=settings.output_dir / "discover_cache.json")
    added, report = [], []
    for f in failures:
        prof = by_email.get(f["email"].lower())
        if prof is None:
            continue
        tag = f"Replacement for {prof.name}"
        if any(tag in n for n in already):
            report.append({"for": prof.name, "status": "already replaced"})
            continue
        found = 0
        while found < limit:
            cand, why = find_colleague(disc, prof, skip_ids, skip_names)
            if cand is None:
                report.append({"for": prof.name, "status": why})
                break
            homepage = disc.ror_domains(cand.ror)[1]
            new = to_professor(cand, homepage)
            new.notes = (f"{tag} ({f['root_cause']}, {f['email']}); same institution, "
                         f"topic: {why[:60]}. " + new.notes)
            added.append(new)
            skip_ids.add(cand.author_id)
            skip_names.add(name_key(cand.name))
            report.append({"for": prof.name, "replacement": cand.name,
                           "institution": cand.institution, "h_index": cand.h_index})
            progress(f"  {prof.name[:28]:<28} -> {cand.name} ({cand.institution}, h={cand.h_index})"
                     .encode("ascii", "replace").decode())
            found += 1
    disc.save_cache()
    if added:
        append_rows(path, added)
    return {"marked_invalid": marked, "added": len(added), "details": report}
