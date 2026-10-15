"""Pipeline accounting: how many rows reached each stage, and why the rest didn't."""
from __future__ import annotations

import csv
import json
import shutil
import tempfile
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .excel_io import _column_index, _find_header_row, sidecar_path

SEND_LOG_FIELDS = ["time", "mode", "status", "email", "name", "sheet", "row", "detail"]


def _bucket(status: str) -> str:
    """Collapse 'skipped: local score 0.21 < 0.25' style values into one bucket."""
    status = (status or "").strip()
    if not status:
        return ""
    if status.startswith("skipped: local score"):
        return "skipped: local score below threshold"
    if status.startswith("ok ("):
        return "ok"
    if status.startswith("failed:") and len(status) > 60:
        return status[:60] + "..."
    return status


def _open_readonly(path: Path):
    """Open a workbook even while Excel holds it, via a temp copy."""
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except OSError:
        tmp = Path(tempfile.gettempdir()) / f"outreach-stats-{path.name}"
        shutil.copyfile(path, tmp)
        return load_workbook(tmp, read_only=True, data_only=True)


def collect(path: str | Path, sheet: str | None = None) -> dict:
    path = Path(path)
    wb = _open_readonly(path)
    pending = {}
    side = sidecar_path(path)
    if side.exists():
        try:
            pending = {(p["sheet"], p["row"]): p["status"]
                       for p in json.loads(side.read_text(encoding="utf-8"))}
        except (OSError, ValueError, KeyError):
            pending = {}

    c = Counter()
    llm, send, reply, email = Counter(), Counter(), Counter(), Counter()
    sent_days = Counter()
    sheets = [wb[sheet]] if sheet else list(wb.worksheets)
    for ws in sheets:
        # read_only worksheets don't support random cell access well; load rows once.
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header_row = _find_header_row(ws)
        headers = {k: v - 1 for k, v in _column_index(ws, header_row).items()}

        for r_idx, row in enumerate(rows[header_row:], start=header_row + 1):
            def get(name):
                i = headers.get(name.lower())
                v = row[i] if i is not None and i < len(row) else None
                return str(v).strip() if v is not None else ""

            name, addr = get("Professor Name") or get("Name"), get("Email Address") or get("Email")
            if not (name or addr):
                continue
            c["total"] += 1
            if addr:
                c["with_email"] += 1
            if get("Email Status"):
                email[get("Email Status").lower()] += 1
            if get("LLM Status"):
                llm[_bucket(get("LLM Status"))] += 1
            if get("Custom Email Draft"):
                c["drafted"] += 1
            status = pending.get((ws.title, r_idx)) or get("Send Status")
            if status:
                send[_bucket(status)] += 1
            elif get("Custom Email Draft"):
                send["not attempted yet"] += 1
            if status.lower() == "sent":
                day = get("Sent At")[:10]
                if day:
                    sent_days[day] += 1
            if get("Reply Status"):
                reply[get("Reply Status").split()[0]] += 1

    c["sent"] = send.get("sent", 0)
    c["review_sent"] = send.get("review-sent", 0)
    return {"workbook": str(path), "counts": c, "email": email, "llm": llm,
            "send": send, "reply": reply, "sent_days": sent_days,
            "pending_sidecar": len(pending)}


def append_send_log(log_path: str | Path, mode: str, rows: list[dict],
                    updates: list[dict]) -> None:
    """Append one line per send attempt so history survives workbook edits."""
    if not updates:
        return
    log_path = Path(log_path)
    by_key = {(r["sheet"], r["row"]): r for r in rows}
    new = not log_path.exists()
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with log_path.open("a", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=SEND_LOG_FIELDS)
        if new:
            w.writeheader()
        for u in updates:
            src = by_key.get((u.get("sheet"), u.get("row")), {})
            w.writerow({
                "time": u.get("sent_at") or now, "mode": mode,
                "status": u.get("status", ""), "email": src.get("email", ""),
                "name": src.get("name", ""), "sheet": u.get("sheet", ""),
                "row": u.get("row", ""), "detail": str(u.get("detail", ""))[:200],
            })


def read_send_log(log_path: str | Path) -> Counter:
    """Count logged attempts per (mode, status)."""
    log_path = Path(log_path)
    out = Counter()
    if not log_path.exists():
        return out
    with log_path.open(newline="", encoding="utf-8") as fh:
        for rec in csv.DictReader(fh):
            out[(rec.get("mode", ""), _bucket(rec.get("status", "")))] += 1
    return out
