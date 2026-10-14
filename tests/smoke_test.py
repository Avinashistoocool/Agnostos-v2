"""End-to-end smoke test with a mocked OpenAI client (no API calls, no emails sent)."""
from __future__ import annotations

import json
import shutil
import sys
from dataclasses import replace
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from outreach.config import settings
from outreach.excel_io import (load_for_sending, load_targets, mark_sent,
                               update_workbook, write_new_workbook)
from outreach.llm import LLMEvaluator
from outreach.mailer import Mailer
from outreach.matcher import filter_candidates
from outreach.resume_parser import parse_resume
from outreach.scraper import ScrapeResult, _html_to_text, _looks_blocked, extract_professors

import _paths

settings.resume_path = _paths.resume_path(settings)
settings.target_xlsx = _paths.target_xlsx(settings)
MIN_TARGETS = 10 if _paths.using_fixtures(settings) else 40

TMP = Path(__file__).resolve().parent / "_tmp"


class _PromptDetails:
    cached_tokens = 1000


class _Usage:
    prompt_tokens = 1420
    completion_tokens = 310
    prompt_tokens_details = _PromptDetails()


class _Choice:
    finish_reason = "stop"

    def __init__(self, content):
        self.message = type("Msg", (), {"content": content})()


class _Resp:
    def __init__(self, text):
        self.choices = [_Choice(text)]
        self.usage = _Usage()


class FakeCompletions:
    def __init__(self):
        self.calls = 0

    def create(self, **kwargs):
        self.calls += 1
        assert kwargs["response_format"]["type"] == "json_schema"
        assert kwargs["messages"][0]["role"] == "system"
        ctx = kwargs["messages"][1]["content"]
        name = ctx.split("Professor: ")[1].split("\n")[0]
        return _Resp(json.dumps({
            "probability": 40 + self.calls,
            "justification": f"Strong overlap on jet tagging with {name}.",
            "subject": "Summer internship enquiry - jet tagging",
            "body": f"Dear {name},\n\nI read your recent work.\n\nBest,\nAlex",
        }))


class FakeClient:
    def __init__(self):
        self.chat = type("Chat", (), {"completions": FakeCompletions()})()


def _strip_result_columns(path: Path) -> None:
    """Remove any columns this tool previously wrote, so the fixture is pristine."""
    import openpyxl
    from outreach.excel_io import RESULT_COLUMNS, _find_header_row

    wanted = {c.lower() for c in RESULT_COLUMNS}
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        header_row = _find_header_row(ws)
        doomed = [c for c in range(ws.max_column, 0, -1)
                  if str(ws.cell(row=header_row, column=c).value or "").strip().lower() in wanted]
        for col in doomed:
            ws.delete_cols(col)
    wb.save(path)


def main() -> int:
    if TMP.exists():
        shutil.rmtree(TMP)
    TMP.mkdir(parents=True)

    print("[1] resume parse")
    profile = parse_resume(settings.resume_path)
    assert profile.name and profile.email and profile.projects, "resume parse failed"
    print(f"    ok: {profile.name}, {len(profile.projects)} projects")

    print("[2] excel ingest")
    src = TMP / "targets.xlsx"
    shutil.copy2(settings.target_xlsx, src)
    _strip_result_columns(src)  # the live workbook may already hold real drafts
    professors = load_targets(src)
    assert len(professors) >= MIN_TARGETS, len(professors)
    print(f"    ok: {len(professors)} professors")

    print("[3] local filter")
    selected, rejected = filter_candidates(professors, profile, 0.25, 8)
    assert len(selected) == 8 and rejected
    assert all(p.local_score >= 0.25 for p in selected)
    print(f"    ok: {len(selected)} selected, {len(rejected)} rejected, "
          f"top={selected[0].local_score}")

    print("[3b] thin-context professors are rejected before the LLM")
    from outreach.models import Professor as _P
    thin = [_P(name="Dr Ghost", email="g@x.edu", institution="X",
               research_area="CMS 511b")]
    sel_thin, rej_thin = filter_candidates(thin + professors[:5], profile, 0.0, 10)
    assert any("too little research context" in p.llm_status for p in rej_thin), \
        [p.llm_status for p in rej_thin]
    assert all(p.name != "Dr Ghost" for p in sel_thin)
    print("    ok: directory stubs filtered out")

    print("[4] estimation gate maths")
    ev = LLMEvaluator(settings, profile)
    est = ev.estimate(selected)
    assert est.n_calls == 8
    if settings.is_free_tier:
        assert est.cost_usd == 0 and est.free_tier
        print(f"    ok: {est.n_calls} calls, free tier ({settings.provider})")
    else:
        assert est.cost_usd > 0
        print(f"    ok: ${est.cost_usd:.4f} for {est.n_calls} calls")

    print("[4b] paid-provider pricing still computes")
    paid = replace(settings, provider="openai", draft_model="gpt-4.1")
    est_paid = LLMEvaluator(paid, profile).estimate(selected)
    assert est_paid.cost_usd > 0 and not est_paid.free_tier
    print(f"    ok: ${est_paid.cost_usd:.4f} on openai/gpt-4.1")

    print("[5] mocked LLM evaluation")
    ev._client = FakeClient()
    ev.evaluate_all(selected)
    assert all(p.probability > 0 for p in selected), [p.llm_status for p in selected]
    assert all(p.email_body.startswith("Dear ") for p in selected)
    assert all("\n" in p.email_body for p in selected), "escaped newlines not decoded"
    assert ev.actual_cost() > 0 or settings.is_free_tier
    print(f"    ok: {ev.usage_report()}")

    print("[5b] abort after repeated failures")
    class BrokenCompletions:
        def create(self, **_kw):
            raise ValueError("bad gateway")

    class BrokenClient:
        chat = type("Chat", (), {"completions": BrokenCompletions()})()

    import copy
    victims = [copy.deepcopy(p) for p in professors if p.email][:6]
    for v in victims:
        v.probability, v.llm_status = 0, ""
    ev2 = LLMEvaluator(settings, profile)
    ev2._client = BrokenClient()
    ev2.evaluate_all(victims)
    assert ev2.usage["calls"] == 0
    assert any("aborted after repeated failures" in v.llm_status for v in victims), \
        [v.llm_status for v in victims]
    print("    ok: stopped early instead of burning the whole list")

    print("[5c] draft body formatting repair")
    from outreach.llm import tidy_body
    wall = ("I read your work on jet tagging. I built GDS-ParT in PyTorch. "
            "I am available May to July 2027. I can apply via DAAD WISE. "
            "Would you have capacity to host? May I send a proposal?")
    fixed = tidy_body(wall, "Prof. Michael Spannowsky")
    assert fixed.startswith("Dear Prof. Spannowsky,\n\n"), fixed[:40]
    assert fixed.count("\n\n") >= 2, fixed
    assert tidy_body("Body.", "Prof. Dr. Anna Schmidt").startswith("Dear Prof. Schmidt,")
    assert tidy_body("Body.", "").startswith("Dear Professor,")
    good = "Dear Dr. Butter,\n\nOne.\n\nTwo.\n\nBest,\nAlex"
    assert tidy_body(good, "Dr. Butter") == good, "mangled a well-formed body"
    print("    ok: salutation added, wall split, well-formed body untouched")

    print("[5d] humanizing layer")
    from outreach.humanize import deslop, has_hard_claims
    from outreach.humanize import score as ai_score
    sloppy = ("Dear Prof. X,\n\nI am writing to express my interest. I hope this email "
              "finds you well. Moreover, I utilised cutting-edge methods and my model "
              "achieved superior results.\n\nBest,\nAlex")
    before, _ = ai_score(sloppy)
    cleaned = deslop(sloppy)
    after, _ = ai_score(cleaned)
    assert after < before, (before, after)
    for banned in ("I am writing to", "utilised", "cutting-edge", "Moreover",
                   "achieved", "superior", "finds you well"):
        assert banned not in cleaned, banned
    assert cleaned.startswith("Dear Prof. X,")
    assert has_hard_claims("improved the accuracy by 12%")
    assert not has_hard_claims("I built a transformer variant and ran it on a subset.")
    print(f"    ok: AI-tell score {before} -> {after}, claims detected")

    print("[6] write back to source workbook (Mode B)")
    update_workbook(src, selected + rejected, make_backup=True)
    backups = list(TMP.glob("targets.backup-*.xlsx"))
    assert backups, "no backup created"
    rows = load_for_sending(src)
    assert len(rows) == 8, len(rows)
    assert rows[0]["subject"] and rows[0]["probability"] > 0
    print(f"    ok: {len(rows)} sendable rows, backup={backups[0].name}")

    print("[6b] idempotent re-write does not duplicate columns")
    import openpyxl
    sheet = openpyxl.load_workbook(src).sheetnames[0]
    before = openpyxl.load_workbook(src)[sheet].max_column
    update_workbook(src, selected + rejected, make_backup=False)
    after = openpyxl.load_workbook(src)[sheet].max_column
    assert before == after, (before, after)
    print(f"    ok: column count stable at {after}")

    print("[7] new workbook (Mode A output)")
    out = TMP / "new.xlsx"
    write_new_workbook(out, selected)
    new_rows = load_for_sending(out)
    assert len(new_rows) == 8
    assert new_rows[0]["probability"] >= new_rows[-1]["probability"], "not sorted"
    print(f"    ok: {len(new_rows)} rows written and re-read")

    print("[8] mailer dry-run + attachments")
    settings.attachments = [str(settings.resume_path)]
    mailer = Mailer(settings, profile, dry_run=True)
    updates = mailer.send_batch(rows)
    assert len(updates) == 8 and all(u["status"] == "dry-run" for u in updates)
    msg = mailer.send_one(rows[0])["message"]
    body = msg.get_body(("plain",)).get_content()
    assert profile.email in body and profile.name in body, "signature missing"
    atts = [a.get_filename() for a in msg.iter_attachments()]
    assert atts == [Path(settings.resume_path).name], atts
    print(f"    ok: 8 dry-run msgs, attachment={atts[0]}")

    print("[9] send-status writeback + resume safety")
    mark_sent(src, updates)
    rows2 = load_for_sending(src)
    assert all(r["send_status"] == "dry-run" for r in rows2)
    # A dry run must NOT block a later real send.
    still_pending = Mailer(settings, profile, dry_run=True).send_batch(rows2)
    assert len(still_pending) == 8, "dry-run wrongly treated as already sent"
    # A genuine send must block a repeat.
    sent_updates = [{**u, "status": "sent"} for u in updates]
    mark_sent(src, sent_updates)
    rows3 = load_for_sending(src)
    again = Mailer(settings, profile, dry_run=True).send_batch(rows3)
    assert not again, "already-sent rows were not skipped"
    forced = Mailer(settings, profile, dry_run=True).send_batch(rows3, resend=True)
    assert len(forced) == 8
    print("    ok: dry-run never blocks, real sends do, --resend overrides")

    print("[9a] a locked workbook parks statuses instead of losing them")
    import outreach.excel_io as _xio
    from outreach.excel_io import is_writable, sidecar_path
    assert is_writable(src)
    locked_updates = [{**u, "status": "failed: SMTPServerDisconnected",
                       "sent_at": ""} for u in updates[:3]]
    real_load = _xio.load_workbook

    def _locked(*a, **kw):
        raise OSError("file is open in Excel")

    _xio.load_workbook = _locked
    try:
        deferred = mark_sent(src, locked_updates)
    finally:
        _xio.load_workbook = real_load
    assert deferred == sidecar_path(src) and deferred.exists()
    by_row = {r["row"]: r for r in load_for_sending(src)}
    assert all(by_row[u["row"]]["send_status"] == "failed: SMTPServerDisconnected"
               for u in locked_updates), "parked statuses were not applied on read"
    # The next successful write drains the sidecar into the workbook.
    assert mark_sent(src, []) is None
    assert not sidecar_path(src).exists()
    by_row = {r["row"]: r for r in load_for_sending(src)}
    assert all(by_row[u["row"]]["send_status"] == "failed: SMTPServerDisconnected"
               for u in locked_updates), "parked statuses were not persisted"
    mark_sent(src, [{**u, "status": "sent"} for u in locked_updates])
    print("    ok: parked to sidecar, honoured on read, drained on next write")

    print("[9b] review redirect keeps the real target visible")
    reviewer = Mailer(settings, profile, dry_run=True, redirect_to="me@example.com")
    msg = reviewer.send_one(rows3[0])["message"]
    assert msg["To"] == "me@example.com", msg["To"]
    assert rows3[0]["email"] in msg["Subject"], msg["Subject"]
    print(f"    ok: {msg['Subject'][:58]}")

    print("[9c] .eml export opens in a mail client")
    eml_dir = TMP / "drafts"
    written = Mailer(settings, profile, dry_run=True).save_eml(rows3[0], eml_dir)
    assert written.exists() and written.suffix == ".eml"
    import email as _email
    from email import policy as _policy
    parsed = _email.message_from_bytes(written.read_bytes(), policy=_policy.default)
    assert parsed["To"] and parsed["Subject"]
    assert [a.get_filename() for a in parsed.iter_attachments()] == \
        [Path(settings.resume_path).name]
    print(f"    ok: {written.name}")

    print("[9d] transient SMTP drop reconnects; persistent failure aborts")
    import smtplib as _smtplib

    class _FlakyServer:
        def __init__(self, fail_times):
            self.fail_times = fail_times
            self.delivered = 0

        def send_message(self, _msg):
            if self.fail_times > 0:
                self.fail_times -= 1
                raise _smtplib.SMTPServerDisconnected("connection reset")
            self.delivered += 1

    flaky = Mailer(settings, profile)
    flaky.dry_run = False
    server = _FlakyServer(fail_times=1)
    flaky._server = server
    flaky.connect = lambda: setattr(flaky, "_server", server)
    assert flaky.send_one(rows3[0])["status"] == "sent", "transient drop was not retried"
    assert server.delivered == 1

    dead = Mailer(settings, profile)
    dead.dry_run = False
    dead_server = _FlakyServer(fail_times=999)
    dead._server = dead_server
    dead.connect = lambda: setattr(dead, "_server", dead_server)
    dead.s = replace(settings, send_delay_min=0, send_delay_max=0)
    aborts = []
    dead_updates = dead.send_batch(
        rows3, resend=True,
        progress=lambda *a: aborts.append(a[3]["status"]))
    assert "aborted" in aborts, aborts
    assert len(dead_updates) < len(rows3), (len(dead_updates), len(rows3))
    print(f"    ok: retried once, then stopped after {len(dead_updates)} failures")

    print("[10] scraper extraction heuristics")
    html = """<html><body>
      <div class="person"><h3>Prof. Dr. Anna Schmidt</h3>
      <p>Machine learning for collider physics and jet tagging.</p>
      <p>Email: anna.schmidt [at] uni-example [dot] de</p></div>
      <div class="person"><h3>Dr. Marco Rossi</h3>
      <p>Graph neural networks for tracking.</p>
      <a href="mailto:m.rossi@uni-example.de">m.rossi@uni-example.de</a></div>
      <p>General enquiries: info@uni-example.de</p>
      <img src="logo@2x.png">
    </body></html>"""
    res = ScrapeResult("https://uni-example.de/staff", html, _html_to_text(html))
    found = extract_professors(res, "Uni Example", "Germany", "Physics")
    emails = sorted(p.email for p in found if p.email)
    assert emails == ["anna.schmidt@uni-example.de", "m.rossi@uni-example.de"], emails
    assert all(p.institution == "Uni Example" for p in found)
    print(f"    ok: extracted {emails} (generic info@ dropped)")

    print("[10b] emails behind profile links + role-suffix cleanup")
    html2 = """<html><body><table>
      <tr><td><a href="/people/j.smith">Prof Jane Smith</a></td><td>Professor</td></tr>
      <tr><td><a href="/people/k.lee">Dr Kai Lee</a></td><td>Senior Lecturer</td></tr>
    </table></body></html>"""
    text2 = "Prof Jane Smith\tProfessor\nDr Kai Lee\tSenior Lecturer"
    found2 = extract_professors(ScrapeResult("https://u.ac.uk/dept/", html2, text2),
                                "U", "UK", "Physics")
    names = sorted(p.name for p in found2)
    assert names == ["Dr Kai Lee", "Prof Jane Smith"], names
    urls = sorted(p.profile_url for p in found2)
    assert urls == ["https://u.ac.uk/people/j.smith", "https://u.ac.uk/people/k.lee"], urls
    assert all(not p.email and p.notes for p in found2)
    print(f"    ok: {names} resolved to profile pages for the second hop")

    print("[11] blocked-page detection")
    assert _looks_blocked("<html><title>Just a moment...</title>cf-chl</html>")
    assert not _looks_blocked("<html><body>Faculty directory</body></html>")
    # A long, genuine page that merely mentions Cloudflare must not be flagged.
    big = "<html><body>" + ("Faculty directory cloudflare cdn. " * 300) + "</body></html>"
    assert not _looks_blocked(big)
    print("    ok (no false positive on large pages)")

    print("[12] email verification scoring (no network)")
    from outreach.verify import (EmailCheck, check_email, is_academic,
                                 name_match, same_site)
    # Local parts a university actually assigns.
    for nm, em in [("Prof. Michael Spannowsky", "michael.spannowsky@kit.edu"),
                   ("Prof. Max Welling", "m.welling@uva.nl"),
                   ("Prof. Ivo van Vulpen", "i.vanvulpen@uva.nl"),
                   ("Dr. Roberto Ruiz de Austri", "rruiz@ific.uv.es"),
                   ("Prof. Niels Richard Hansen", "Niels.R.Hansen@math.ku.dk")]:
        assert name_match(nm, em) >= 0.75, (nm, em)
    # Someone else's mailbox, and a department office, must not pass.
    assert name_match("Prof. Michael Spannowsky", "gregor.kasieczka@kit.edu") == 0.0
    assert name_match("Prof. Tilman Plehn", "info@uni-heidelberg.de") == 0.0
    # Sub-domains belong to their parent; unrelated universities do not.
    assert same_site("thphys.uni-heidelberg.de", "uni-heidelberg.de")
    assert same_site("kit.edu", "kit.edu")
    assert not same_site("kit.edu", "durham.ac.uk")
    assert not same_site("kit.edu", "")
    assert is_academic("unige.ch") and not is_academic("gmail.com")

    good = check_email("Prof. Michael Spannowsky", "michael.spannowsky@kit.edu",
                       "KIT", check_dns=False)
    assert good.status == "ok" and good.confidence == 100, good
    wrong_person = check_email("Prof. Michael Spannowsky",
                               "gregor.kasieczka@kit.edu", "KIT", check_dns=False)
    assert wrong_person.status == "suspect" and wrong_person.confidence <= 50
    # A domain that disagrees with the one found on their live page is docked,
    # but an institution string that merely reads differently never is.
    assert check_email("Prof. Michael Spannowsky", "michael.spannowsky@kit.edu",
                       "Durham University", check_dns=False).status == "ok"
    stale = check_email("Prof. Michael Spannowsky", "michael.spannowsky@kit.edu",
                        "", ("durham.ac.uk",), check_dns=False)
    assert stale.confidence < good.confidence, stale
    assert "durham.ac.uk" in stale.summary(), stale
    for bad in ("", "not-an-email", "two@@at.com"):
        assert check_email("Someone", bad, "KIT", check_dns=False).status == "invalid"
    assert EmailCheck(status="invalid").sendable is False
    # Addresses published with a harvester trap in the domain are repaired,
    # not discarded - the person wrote a real address under the label.
    from outreach.verify import EmailFinder as _EF
    assert _EF._emails_in("mail nicola.serra@SPAMNOT.cern.ch here") == \
        ["nicola.serra@cern.ch"]
    print("    ok: right person required, stale domains flagged, junk rejected")

    print("[12a] invalid rows are skipped by the mailer")
    rows12 = load_for_sending(src)[:2]
    for r in rows12:
        r["send_status"] = ""
    rows12[0]["email_status"] = "invalid"
    with Mailer(settings, profile, dry_run=True) as m12:
        skipped = m12.send_batch(rows12)
    assert skipped[0]["status"].startswith("skipped: invalid")
    assert skipped[1]["status"] == "dry-run"
    with Mailer(settings, profile, dry_run=True) as m12:
        forced = m12.send_batch(rows12, allow_unverified=True)
    assert all(u["status"] == "dry-run" for u in forced), forced
    print("    ok: blocked by default, --allow-unverified overrides")

    print("[13] topic discovery helpers (no network)")
    from outreach.discover import (Candidate, _best_name, _this_year,
                                   resolve_countries, to_professor)
    from outreach.models import Professor
    assert resolve_countries("uk") == ["gb"]
    assert resolve_countries("de,ch") == ["de", "ch"]
    assert set(resolve_countries("dach")) == {"de", "at", "ch"}
    assert _best_name(["T. Golling", "TOBIAS GOLLING", "Golling, T.",
                       "Tobias Golling"]) == "Tobias Golling"
    assert _best_name(["Tilman Plehn"]) == "Tilman Plehn"
    assert _best_name(["Dr. Tilman Plehn"]) == "Tilman Plehn"
    lead = Candidate(author_id="A1", name="A", hits=4, lead_hits=4,
                     newest_year=_this_year(), h_index=60)
    bulk = Candidate(author_id="A2", name="B", hits=4, lead_hits=0,
                     newest_year=0, h_index=0)
    assert lead.score > bulk.score, (lead.score, bulk.score)
    # Unverified discovered rows must not claim an "Email Confidence" of 0%.
    assert to_professor(lead).email_status == ""
    print("    ok: country groups, name repair, lead-author ranking, blank status")

    print("[14] batch planner, append, recheck helpers (no network)")
    from outreach.batches import name_key, plan_batches, split_queries
    from outreach.excel_io import append_rows, load_targets as _load14
    from outreach.recheck import same_institution
    assert split_queries("a; b | c\nd") == ["a", "b", "c", "d"]
    plan = plan_batches(["x", "y"], 1000, 4)
    assert [b.quota for b in plan] == [250] * 4
    assert [b.query for b in plan] == ["x", "y", "x", "y"]
    assert sum(b.quota for b in plan_batches(["x"], 10, 4)) == 10
    assert name_key("Jürgen Müller") == name_key("Muller, Jurgen")
    book = TMP / "append.xlsx"
    append_rows(book, [Professor(name="A One", institution="EPFL",
                                 email="a@epfl.ch", openalex_id="A1")])
    append_rows(book, [Professor(name="B Two", institution="ETH",
                                 email="b@ethz.ch", openalex_id="A2")])
    rows = _load14(book)
    assert [r.name for r in rows] == ["A One", "B Two"], [r.name for r in rows]
    assert same_institution("Heidelberg University", "Universität Heidelberg")
    assert not same_institution("Heidelberg University", "ETH Zurich")
    print("    ok: even quotas, round-robin topics, name keys, append, institution match")

    print("[15] review_1 regressions (no network)")
    import json as _json
    from outreach.batches import fingerprint
    from outreach.cli import _drop_sent
    from outreach.excel_io import sidecar_path, update_workbook as _upd15
    from outreach import verify as _v
    # Defaults keep the old fingerprint, so in-progress runs still resume.
    base = fingerprint(["q"], 10, 2, ["de"], 2021, 8, 20)
    assert fingerprint(["q"], 10, 2, ["de"], 2021, 8, 20, {}) == base
    assert fingerprint(["q"], 10, 2, ["de"], 2021, 8, 20,
                       {"min_relevance": 0.3}) != base
    # A send parked in the sidecar counts as done.
    side = sidecar_path(book)
    side.write_text(_json.dumps([{"sheet": rows[0].source_sheet,
                                  "row": rows[0].source_row,
                                  "status": "sent", "sent_at": ""}]),
                    encoding="utf-8")
    left = _drop_sent(book, _load14(book))
    assert [r.name for r in left] == ["B Two"], [r.name for r in left]
    side.unlink()
    # update_workbook fills a blank OpenAlex ID but never overwrites one.
    r2 = _load14(book)
    r2[0].openalex_id, r2[1].openalex_id = "CHANGED", "A2"
    _upd15(book, r2, make_backup=False)
    assert [r.openalex_id for r in _load14(book)] == ["A1", "A2"]
    # --apply keeps a verified address and only replaces a bad/missing one.
    orig = _v.verify_professor

    def fake(prof, **kw):
        if prof.email == "good@epfl.ch":
            return _v.EmailCheck(email=prof.email, status=_v.OK, confidence=90,
                                 suggestion="g.alt@epfl.ch", source="page")
        if prof.email == "new@epfl.ch":
            return _v.EmailCheck(email=prof.email, status=_v.OK, confidence=95)
        return _v.EmailCheck(email=prof.email, status=_v.SUSPECT, confidence=30,
                             suggestion="new@epfl.ch", source="page")
    _v.verify_professor = fake
    try:
        pa = Professor(name="Keep Me", email="good@epfl.ch")
        pb = Professor(name="Fix Me", email="bad@old.org")
        _v.verify_many([pa, pb], workers=2, repair=True, apply=True, check_dns=False)
    finally:
        _v.verify_professor = orig
    assert pa.email == "good@epfl.ch" and "g.alt@epfl.ch" in pa.email_suggestion
    assert pb.email == "new@epfl.ch" and "replaced bad@old.org" in pb.email_notes
    with _v._quiet_stderr():
        with _v._quiet_stderr():
            pass
    assert _v._QUIET["depth"] == 0
    print("    ok: fingerprint compat, sidecar skip, fill-blank ids, safe --apply")

    shutil.rmtree(TMP)
    print("\nALL SMOKE TESTS PASSED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
